"""Create a dated, explicitly documented merchant presentation workbook offline."""

import argparse
import csv
import json
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]


def create_pack(output: Path, cutoff: date):
    source = json.loads((ROOT / "ml/examples/walmart-scenario-history.json").read_text(encoding="utf-8"))
    sample = next(s for s in source["stores"] if s["id"] == "walmart-demo-4")
    delta = cutoff - date.fromisoformat(source["source_end"])
    labels = {
        9: ("버터 소금빵", "베이커리", 2800, 1500, 3, 48),
        27: ("햄치즈 샌드위치", "간편식", 4800, 2600, 3, 8),
        90: ("과일 컵", "신선식품", 4500, 2700, 3, 2),
        93: ("그릭 요거트", "유제품", 3900, 2100, 5, 1),
        96: ("초코 쿠키", "베이커리", 2200, 1000, 7, 0),
        97: ("콜드브루", "음료", 3500, 1800, 7, 0),
    }
    book = Workbook()
    guide = book.active
    guide.title = "안내"
    guide.append(["항목", "설명"])
    guide.append(["매장명 예시", "화개마켓 댈러스점"])
    guide.append(["지역", "댈러스, 텍사스 / America/Chicago"])
    guide.append(["마감 기준일", cutoff.isoformat()])
    guide.append(["자료 출처", "Walmart Stormy Weather 공개 판매 112일의 날짜 이동. 상품명·분류·공급사·가격(원)·재고는 더미 가정값이며 원본 상품을 식별한 명칭이 아닙니다."])
    guide.append(["날씨", "원본 매장 좌표는 미공개. 댈러스는 선택한 날씨 위치로 가정합니다. 과거 원본 관측과 현재 댈러스 예보를 잇는 실험입니다."])
    guide.append(["모델", "기존 Walmart TFT 체크포인트 사용. 예보 입력 성능은 별도 검증 전이며 기존 persistence 평가 지표를 적용하지 않습니다."])
    guide.append(["입력 기준", "당일 판매는 마감한 것으로 가정. 재고 시트의 미래 입고도 직접 확인한 뒤 저장하세요."])
    products = book.create_sheet("상품")
    products.append(["id", "name", "category", "supplier", "price", "cost", "lead_days", "shelf_days", "review_days"])
    stocks = book.create_sheet("재고")
    stocks.append(["product_id", "measured_on", "quantity", "expires_on", "arrives_on"])
    for p in sample["products"]:
        name, category, price, cost, shelf, quantity = labels[int(p["model_item_nbr"])]
        products.append([p["id"], name, category, "화개유통", price, cost, 1, shelf, 2])
        stocks.append([p["id"], cutoff.isoformat(), quantity, (cutoff + timedelta(days=shelf)).isoformat(), None])
        if quantity:
            stocks.append([p["id"], cutoff.isoformat(), max(1, quantity // 3), (cutoff + timedelta(days=shelf + 1)).isoformat(), (cutoff + timedelta(days=1)).isoformat()])
    sales = book.create_sheet("판매")
    sales.append(["date", "product_id", "units_sold"])
    sale_rows = [[(date.fromisoformat(s["date"]) + delta).isoformat(), s["product_id"], s["units_sold"]] for s in sample["sales"]]
    for row in sale_rows:
        sales.append(row)
    model = book.create_sheet("모델연결")
    model.append(["profile", "anchor_date"])
    model.append(["walmart-demo-4", cutoff.isoformat()])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=12)
            cell.fill = PatternFill("solid", fgColor="1859A9")
        sheet.row_dimensions[1].height = 28
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(18, max(len(str(c.value or "")) for c in column) + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11, color="243149")
                cell.alignment = Alignment(vertical="center", wrap_text=sheet.title == "안내")
    guide.column_dimensions["B"].width = 100
    for i in range(2, guide.max_row + 1):
        guide.row_dimensions[i].height = 45
    output.mkdir(parents=True, exist_ok=True)
    path = output / f"화개마켓_매장자료_{cutoff}.xlsx"
    book.save(path)
    with (output / f"판매기록_{cutoff}.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["date", "product_id", "units_sold"])
        writer.writerows(sale_rows)
    (output / "사용순서.txt").write_text(f"""화개장날 · 직접 자료를 가져와 사용하는 순서

1. https://hwagae.milky.kr/ 에 기존 점주 계정으로 로그인합니다.
2. 매장 선택 옆 + → 매장명 '화개마켓 댈러스점' → 지역 '댈러스, 텍사스' → 매장 추가.
3. '우리 매장 자료를 가져오세요'에서 화개마켓_매장자료_{cutoff}.xlsx 선택.
4. 내용 확인하기 → 상품 6개, 판매 672건, 재고 6개 상품을 확인합니다.
5. 현재 재고·입고 예정 확인에 체크 → 확인한 자료 저장.
6. 홈에서 예상 판매와 재고 현황 확인 → 7일 판매 예측에서 실제 미국 날씨 확인.
7. 발주 계획 → 발주 예정일을 이틀 뒤로 선택 → 발주량 계산.
8. 원하는 수량 수정 → 발주안 저장 → 점주 확정 → CSV 또는 Excel 다운로드.
9. 재고 확인에서 수량을 바꾸고 다시 계산하여 발주 변화도 확인할 수 있습니다.

기준일: 미국 댈러스 현지 {cutoff}. 당일 마감 자료로 가정합니다.
상품명·공급사·가격(원)·재고는 더미 값. 판매 수량은 Walmart 공개 기록 112일의 날짜를 이동했습니다.
표시 상품명은 원본 Walmart 상품의 실제 이름이 아닙니다. 실제 매장 매출·성능을 입증하는 자료가 아닙니다.
원본 매장 좌표는 알려져 있지 않으며 댈러스는 시연에서 선택한 날씨 위치입니다.
기존 Walmart TFT와 연결되며, 예보 입력에 대한 정확도는 별도로 평가해야 합니다.

판매 CSV는 별도 확인용입니다. 통합 엑셀을 저장했다면 같은 CSV를 다시 저장하지 않습니다.
자료를 등록한 후에는 판매 기록 메뉴에서 새 날짜의 CSV를 추가할 수 있습니다.
다른 날짜에 시연하려면 현재 현지 날짜로 파일을 다시 생성하여 새 매장에 가져오세요.
생성 명령: uv run python scripts/make-merchant-data.py --output <저장폴더>
""", encoding="utf-8-sig")
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--as-of", type=date.fromisoformat, default=None)
    args = parser.parse_args()
    print(create_pack(args.output, args.as_of or datetime.now(ZoneInfo("America/Chicago")).date()))
