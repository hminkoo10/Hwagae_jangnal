"""디미고 매점 POS 판매·구매(발주) 원본을 모델 학습 70%·평가 30%·시연 업로드 자료로 변환한다.

원본과 출력은 실제 매장 자료이므로 저장소 밖 폴더에서만 다룬다.
"""

import argparse
import csv
import hashlib
import json
import re
import shutil
import urllib.parse
import urllib.request
from collections import Counter, defaultdict, deque
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# 한국디지털미디어고등학교 (Photon 검색 결과, 앱의 매장 지역 검색과 같은 제공처)
LATITUDE, LONGITUDE, TIMEZONE = 37.3416467, 126.8310305, "Asia/Seoul"
TRAIN_SHARE = 0.7
# 시연 매장의 상품 설정 가정. 발주내역에서 추출한 값이 아니다.
LEAD_DAYS, REVIEW_DAYS = 1, 3
# 거래처는 가상 유통사 10곳. 같은 제조사(품목명 앞 "농심)", "롯 데)" 표기) 상품은 같은 곳에서, 표기가 없으면 종류별로 받는다.
DISTRIBUTORS = ["한솔유통", "대성상사", "미래식품유통", "동방물산", "새한유통", "우리식자재", "성원상사", "태광유통", "금강식품", "영진물류"]
MAKERS = {"롯": "롯데", "롯데제과": "롯데", "해": "해태", "헤태": "해태", "해태제과": "해태", "오": "오리온", "농": "농심",
          "그라운": "크라운", "cj": "CJ"}


def maker_group(order_names, kind):
    makers = Counter(MAKERS.get(key, key) for name in order_names
                     if (m := re.match(r"\s*([^)]{1,12})\)", name)) and (key := m.group(1).replace(" ", "").lstrip(",*")))
    return makers.most_common(1)[0][0] if makers else kind


def assign_distributors(groups):
    """Give each maker group one distributor, largest groups first to the least-loaded one."""
    counts, load, result = Counter(groups.values()), [0] * len(DISTRIBUTORS), {}
    for group in sorted(counts, key=lambda g: (-counts[g], g)):
        index = min(range(len(DISTRIBUTORS)), key=lambda i: (load[i], i))
        result[group], load[index] = DISTRIBUTORS[index], load[index] + counts[group]
    return {pid: result[group] for pid, group in groups.items()}


SHELF_DAYS = {"아이스크림": 365, "음료": 180, "유제품": 30, "사탕·젤리": 365, "과자": 180, "빵·디저트": 60, "간편식": 60, "기타": 180}
# 발주 합계가 판매 합계와 이 범위로 맞는 상품만 재고를 추정한다.
COVERAGE = (0.85, 1.25)
COST_RATIO = (0.3, 0.95)

# 같은 상품이 POS에서 이름만 바뀐 경우. 판매 날짜가 겹치면 합치지 않는다.
SALE_RENAMES = {"제주 삼다수500ml": "제주 삼다수생수500ml"}
# 구매현황의 거래처 규격명 → 판매현황 상품명. 자동 비교로 찾지 못한 확실한 짝만 적는다.
ORDER_ALIASES = {
    "빙그레) 빵또아": "부드러운 빵또아",
    "수입)가메즈에낙/스파이시(매운맛)14g*30*400": "가메즈 에낙 매운맛",
    "농심)추파춥스젤리(사워게코)90g*12入*4": "츄파춥스 젤리 사워 게코",
    "웰그린)애플망고녹차340ml": "웰그린 애플망고",
    "웰그린)자두녹차340ml*20": "웰그린 자두녹차",
    "웰그린)복숭아녹차340ml": "웰그린 복숭아녹차",
    "광동)제주삼다수500mL*20": "제주 삼다수생수500ml",
    "빙그레) 요맘때복숭아": "요맘때 복숭아바",
    "빙그레) 요맘때(블루베리)": "요맘때 블루베리 바",
    "빙그레) 요맘때(슈팅)": "요맘때 슈팅스타콘",
    "빙그레) 요맘때 딸기": "요맘때 딸기바",
    "롯 데) 초코퍼지 바": "초코퍼지",
    "빙그레) 쿠앤크바": "빙그레 쿠앤크바",
    "빙그레) 쿠앤크 말차": "쿠앤크말차바",
    "빙그레) 생귤탱귤 바": "생귤탱귤 감귤",
    "해태음료)코코팜(피치핑크)340ml": "코코팜피치핑크 복숭아340ml",
    "코카)파워에이드(오리지날/pet)600mL*20": "파워에이드 오리지널600ml",
    "빙그레) 붕어싸만코": "붕어싸만코 떡",
    "메로나 오리지널": "메로나",
    "빙그레) 슈퍼콘(쿠엔크)": "슈퍼콘 쿠앤크",
    "동원)보성홍차아이스티(복숭아/제로)500ml*24": "보성홍차 아이스티 제로 복숭아",
    "매일)허쉬초콜릿드링크(쿠키앤크림/팩)235ml*24": "허쉬초콜릿 쿠키앤크림드링크235ml",
    "매일)허쉬초콜릿드링크(오리지날/팩)235ml*24": "허쉬 초콜릿드링크235ml",
    "해 태) 부라보(바니라)": "부라보 콘 바닐라",
    "해 태) 부라보(초코)": "부라보 콘 초코",
    "롯 데) 설레임밀크": "설레임 밀크쉐이크",
}

CATEGORY_RULES = [
    ("아이스크림", r"싸만코|빵또아|메로나|설레임|빠삐코|누가바|쌍쌍바|돼지바|옥동자|요맘때|초코퍼지|찰떡아이스|비비빅|투게더|구구콘|폴라포|탱크보이|더위사냥|수박바|죠스바|스크류바|월드콘|슈퍼콘|부라보|생귤탱귤|아이스주|초코리치바|쿠앤크|메가톤바|찰옥수수|빙빙바|구슬아이스|제로바|빵빠레|모나카|군고구마바|아이스크림|따옴바|^와 |(?<!깔)(?<!팝)콘$"),
    ("유제품", r"우유|요구르트|요거트|치즈|에몽"),
    ("음료", r"ml|드링크|음료|주스|에이드|생수|삼다수|아이스티|커피|라떼|제티|보리|녹차|홍차|밀크티|데자와|비타500|코코팜|파워오투|포카리|사이다|콜라|환타|박카스|립톤|피크닉|이온|워터|탄산|소다|수정과|식혜|토레타|미닛메이드|모구모구|뽀로로|오라떼|솔의눈|봉봉|매실|유자|홍삼|오미자|수염차|우엉차|누룽지|웰그린|티즐|우롱"),
    ("사탕·젤리", r"사탕|젤리|츄파춥스|하리보|마이쮸|캔디|껌|츄잉|구미|곤약|핵짱셔|새콤달콤|카라멜"),
    ("간편식", r"닭가슴살|소시지|맥스봉|천하장사|컵밥|라면|삼각|김밥|햄|도시락|만두"),
    ("빵·디저트", r"빵|케이크|파이|브라우니|카스타드|약과|양갱|떡|와플|마들렌|머핀"),
    ("과자", r"칩|깡|스낵|과자|쿠키|크래커|뻥튀기|감자|팝콘|킥|비쵸비|칸쵸|빈츠|에낙|뿌셔뿌셔|쫄병|꼬깔콘|도리토스|프링글스|오징어|땅콩|아몬드|초코|초콜릿|초코릿|엠앤엠즈|가나|크런키|마가렛트|오레오|칙촉|미쯔|포스틱|닭다리|꽃게랑|고래밥|치킨팝|뻥이요|키커"),
]


def key(name):
    text = name.lower()
    text = re.sub(r"^[^()]*\)", "", text)
    text = re.sub(r"\*.*$", "", text)
    text = re.sub(r"\d+(\.\d+)?\s*(ml|l|g|kg|m|입|入|개)", "", text)
    text = re.sub(r"pet|캔|팩|낱개판매|낱개", "", text)
    return re.sub(r"[^0-9a-z가-힣]", "", text)


def category(name, order_names=()):
    # 구매현황의 아이스크림 거래처 양식은 "롯 데) 돼지바"처럼 접두어 뒤에 공백이 있고 용량이 없다.
    if any(re.match(r"(롯 데|빙그레|해 태|서 주)\) ", n) and not re.search(r"ml", n.lower()) for n in order_names):
        return "아이스크림"
    return next((label for label, pattern in CATEGORY_RULES if re.search(pattern, name.lower())), "기타")


def read_report(path, date_format):
    """판매현황집계·구매현황집계 XLSX: 1행 제목, 2행 머리글, 마지막 총합계·출력시각."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(book.worksheets[0].iter_rows(values_only=True))
    finally:
        book.close()
    if list(rows[1][:6]) != ["일별", "품목별", "수량", "공급가액", "부가세", "합계"]:
        raise ValueError(f"{path.name}: 일별·품목별·수량·공급가액·부가세·합계 열이 필요합니다.")
    records = []
    for row in rows[2:]:
        day, name, quantity, _, _, total = row[:6]
        if day is None or quantity is None or str(day).strip() == "총합계":
            # 구매현황은 발주하지 않은 품목도 수량 없이 나열한다.
            continue
        quantity, total = float(quantity), float(total or 0)
        if quantity != int(quantity):
            raise ValueError(f"{path.name}: 정수가 아닌 수량 {row}")
        records.append((datetime.strptime(str(day).strip(), date_format).date(), str(name).strip(), int(quantity), total))
    return records


def merge_sale_names(sales):
    days = defaultdict(set)
    for day, name, _, _ in sales:
        days[SALE_RENAMES.get(name, name)].add(day)
    for old, new in SALE_RENAMES.items():
        if days.get(old) and days[old] & days[new]:
            raise ValueError(f"판매 날짜가 겹쳐 합칠 수 없습니다: {old} / {new}")
    groups = defaultdict(list)
    for name in days:
        groups[key(name)].append(name)
    canonical, ambiguous = {}, set()
    for group_key, names in groups.items():
        names.sort(key=lambda n: min(days[n]))
        if all(max(days[a]) < min(days[b]) for a, b in zip(names, names[1:])):
            for name in names:
                canonical[name] = names[-1]
        else:
            ambiguous.add(group_key)
            canonical.update({name: name for name in names})
    for old, new in SALE_RENAMES.items():
        canonical[old] = canonical[new]
    return canonical, ambiguous


def daterange(start, end):
    return [start + timedelta(days=i) for i in range((end - start).days + 1)]


def mode_price(rows, end):
    """마지막 판매일까지 최근 28일의 가장 흔한 1개 판매가(부가세 포함)."""
    recent = [(d, q, t) for d, q, t in rows if d <= end and q > 0 and t > 0]
    if not recent:
        return None
    since = max(d for d, _, _ in recent) - timedelta(days=27)
    return Counter(round(t / q) for d, q, t in recent if d >= since).most_common(1)[0][0]


def simulate_stock(receipts, sold, shelf_days, start, end):
    """구매일에 입고, 판매는 선입선출. 기록보다 많이 팔린 수량은 알 수 없는 기초재고로 보고 0에서 멈춘다."""
    batches, open_stock, closing = deque(), {}, {}
    first_receipt = min(receipts) if receipts else None
    for day in daterange(start, end):
        while batches and batches[0][0] + timedelta(days=shelf_days) < day:
            batches.popleft()
        quantity = receipts.get(day, 0)
        if quantity > 0:
            batches.append([day, quantity])
        while quantity < 0 and batches:
            taken = min(-quantity, batches[0][1])
            batches[0][1] -= taken
            quantity += taken
            if not batches[0][1]:
                batches.popleft()
        if first_receipt and day >= first_receipt:
            open_stock[day] = sum(q for _, q in batches)
        remaining = sold.get(day, 0)
        while remaining and batches:
            taken = min(remaining, batches[0][1])
            batches[0][1] -= taken
            remaining -= taken
            if not batches[0][1]:
                batches.popleft()
        closing[day] = [(received, q) for received, q in batches]
    return open_stock, closing


def fetch_weather(start, end):
    params = {"latitude": LATITUDE, "longitude": LONGITUDE, "timezone": TIMEZONE, "temperature_unit": "celsius",
              "precipitation_unit": "mm", "daily": "temperature_2m_max,temperature_2m_min,precipitation_sum"}

    def get(url, extra):
        query = urllib.parse.urlencode({**params, **extra})
        with urllib.request.urlopen(f"{url}?{query}", timeout=30) as response:
            daily = json.load(response)["daily"]
        return {day: (daily["temperature_2m_max"][i], daily["temperature_2m_min"][i], daily["precipitation_sum"][i]) for i, day in enumerate(daily["time"])}

    archive = get("https://archive-api.open-meteo.com/v1/archive", {"start_date": start.isoformat(), "end_date": end.isoformat()})
    rows = {}
    for day in daterange(start, end):
        values = archive.get(day.isoformat())
        if values and None not in values:
            rows[day] = (*values, "Open-Meteo 과거 재분석")
    missing = [day for day in daterange(start, end) if day not in rows]
    if missing:
        recent = get("https://api.open-meteo.com/v1/forecast", {"start_date": min(missing).isoformat(), "end_date": max(missing).isoformat()})
        for day in missing:
            values = recent.get(day.isoformat())
            if values and None not in values:
                rows[day] = (*values, "Open-Meteo 최근 수치모델")
    if len(rows) != len(daterange(start, end)):
        raise ValueError("날씨 자료에 빠진 날짜가 있습니다. 잠시 후 다시 실행하세요.")
    return rows


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def styled(book):
    for sheet in book:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(name="맑은 고딕", color="FFFFFF", bold=True, size=12)
            cell.fill = PatternFill("solid", fgColor="1859A9")
        sheet.row_dimensions[1].height = 28
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(42, max(14, max(len(str(c.value or "")) for c in column) + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11, color="243149")
                cell.alignment = Alignment(vertical="center", wrap_text=sheet.title == "안내")


def create_packs(orders_path, sales_path, output, demo_date):
    sales = read_report(sales_path, "%Y%m%d")
    orders = read_report(orders_path, "%Y/%m/%d")
    canonical, ambiguous = merge_sale_names(sales)
    sales = [(day, canonical[SALE_RENAMES.get(name, name)], quantity, total) for day, name, quantity, total in sales]

    totals = Counter()
    for _, name, quantity, _ in sales:
        totals[name] += quantity
    ids = {name: f"D{rank:03d}" for rank, (name, _) in enumerate(sorted(totals.items(), key=lambda item: (-item[1], item[0])), 1)}
    by_key = {key(name): name for name in ids if key(name) not in ambiguous}
    order_target = {}
    for _, name, _, _ in orders:
        if name in order_target:
            continue
        if name in ORDER_ALIASES:
            order_target[name] = (canonical[ORDER_ALIASES[name]], "직접 대응")
        elif name in canonical:
            order_target[name] = (canonical[name], "같은 이름")
        elif key(name) in by_key:
            order_target[name] = (by_key[key(name)], "규격·접두어 제거 후 같은 이름")
        else:
            order_target[name] = (None, "대응 없음")

    sale_days = sorted({day for day, *_ in sales})
    train_end = sale_days[int(len(sale_days) * TRAIN_SHARE) - 1]
    holdout_start, last_day = train_end + timedelta(days=1), sale_days[-1]
    first_day = sale_days[0]
    # 요일을 유지하도록 7일 단위로만 옮긴다. 시연일과 요일이 다른 마지막 며칠은 업로드에서 뺀다.
    demo_end = last_day - timedelta(days=(last_day.weekday() - demo_date.weekday()) % 7)
    shift = demo_date - demo_end
    if shift.days < 0:
        raise ValueError("시연일은 원본 마지막 판매일 이후여야 합니다.")

    sold = defaultdict(lambda: defaultdict(int))
    price_rows = defaultdict(list)
    for day, name, quantity, total in sales:
        sold[ids[name]][day] += quantity
        price_rows[ids[name]].append((day, quantity, total))
    receipts = defaultdict(lambda: defaultdict(int))
    order_units, order_amount = Counter(), Counter()
    for day, name, quantity, total in orders:
        target = order_target[name][0]
        if target:
            receipts[ids[target]][day] += quantity
            order_units[ids[target]] += quantity
            order_amount[ids[target]] += total
    names = {pid: name for name, pid in ids.items()}
    mapped_orders = defaultdict(list)
    for order_name, (target, _) in order_target.items():
        if target:
            mapped_orders[ids[target]].append(order_name)
    categories = {pid: category(name, mapped_orders[pid]) for pid, name in names.items()}

    profile = {}
    for pid in names:
        price = mode_price(price_rows[pid], demo_end) or 0
        cost = round(order_amount[pid] / order_units[pid]) if order_units[pid] > 0 else None
        coverage = order_units[pid] / max(1, sum(sold[pid].values()))
        reliable = COVERAGE[0] <= coverage <= COVERAGE[1] and cost is not None and price > 0 and COST_RATIO[0] <= cost / price <= COST_RATIO[1]
        profile[pid] = {"price": price, "cost": cost, "coverage": coverage, "reliable": reliable}
    stock = {pid: simulate_stock(receipts[pid], sold[pid], SHELF_DAYS[categories[pid]], first_day, last_day) for pid in names if profile[pid]["reliable"]}

    open_days = set(sale_days)
    weather = fetch_weather(first_day, last_day)

    def panel(start, end, pids=None):
        rows, meta = [], []
        for pid in sorted(names):
            days = [d for d in sold[pid] if start <= d <= end and sold[pid][d] > 0]
            if pids is None and not days:
                continue
            if pids is not None and pid not in pids:
                continue
            first, last = (start, end) if pids is not None else (min(days), max(days))
            opening = stock[pid][0] if pid in stock else {}
            for day in daterange(first, last):
                rows.append([pid, day.isoformat(), sold[pid].get(day, 0), int(day in open_days), opening.get(day, "")])
            price = mode_price(price_rows[pid], last) or ""
            meta.append([pid, names[pid], categories[pid], first.isoformat(), last.isoformat(), price, "발주 누적 기반" if pid in stock else "추정 안 함"])
        return rows, meta

    sales_header = ["product_id", "date", "units_sold", "store_has_sales", "est_stock_open"]
    product_header = ["product_id", "name", "category", "first_date", "last_date", "price", "stock_estimate"]
    weather_header = ["date", "temp_max", "temp_min", "rain_mm", "source"]
    sources = {sales_path.name: sha256(sales_path), orders_path.name: sha256(orders_path)}
    split = {
        "source_files": sources,
        "store": {"name": "디미고사회적협동조합 매점", "latitude": LATITUDE, "longitude": LONGITUDE, "timezone": TIMEZONE},
        "train": {"start": first_day.isoformat(), "end": train_end.isoformat(), "sales_days": sum(d <= train_end for d in sale_days)},
        "holdout": {"start": holdout_start.isoformat(), "end": last_day.isoformat(), "sales_days": sum(d > train_end for d in sale_days)},
        "rule": "판매가 기록된 날짜를 시간순으로 70% / 30% 분할. 섞지 않음.",
    }

    output.mkdir(parents=True, exist_ok=True)
    train_dir, holdout_dir, demo_dir = output / "1_모델담당_70%학습", output / "2_평가보관_30%_공유금지", output / "3_시연업로드"
    for folder in (train_dir, holdout_dir, demo_dir):
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir()

    rows, meta = panel(first_day, train_end)
    write_csv(train_dir / "sales_daily.csv", sales_header, rows)
    write_csv(train_dir / "products.csv", product_header, meta)
    write_csv(train_dir / "weather_daily.csv", weather_header, [[d.isoformat(), *weather[d]] for d in daterange(first_day, train_end)])
    (train_dir / "split.json").write_text(json.dumps({**split, "holdout": {"start": holdout_start.isoformat(), "note": "평가용 30%는 모델 제출 후 공개"}}, ensure_ascii=False, indent=2), encoding="utf-8")
    (train_dir / "README.txt").write_text(f"""디미고 매점 TFT 학습 자료 (앞 70% 기간)

모델 설명: 저장소 docs/model.md

sales_daily.csv
  product_id       상품 코드 (D001 = 1년 판매량 1위). 이름이 바뀐 상품은 하나로 합침
  date             YYYY-MM-DD, 매장 현지(Asia/Seoul)
  units_sold       그날 판매 개수. 상품의 첫·마지막 판매일 사이 판매 행이 없는 날은 실제 0
  store_has_sales  그날 매장 전체에 판매가 1건이라도 있으면 1. 0은 휴무·방학 등 (추론 때는 주어지지 않음)
  est_stock_open   그날 영업 시작 시 추정 재고. 발주 기록이 판매와 맞는 상품만, 첫 입고 이후만. 0이면 품절로 판매가 0일 수 있음
products.csv       이름·종류(이름 규칙으로 분류한 추정)·기간·대표 판매가(원, 부가세 포함)
weather_daily.csv  안산 와동 일별 최고/최저 기온(°C)·강수(mm). 과거 재분석이며 발급 당시 예보가 아님
split.json         원본 해시와 기간

기간: {first_day} ~ {train_end} (판매일 {split['train']['sales_days']}일). 이후 30%는 평가용으로 따로 보관합니다.
원본 POS 자료이므로 Git·공개 저장소에 올리지 마세요.
""", encoding="utf-8")
    shutil.make_archive(str(train_dir), "zip", train_dir)

    rows, meta = panel(first_day, last_day)
    write_csv(holdout_dir / "sales_daily.csv", sales_header, rows)
    write_csv(holdout_dir / "products.csv", product_header, meta)
    write_csv(holdout_dir / "weather_daily.csv", weather_header, [[d.isoformat(), *weather[d]] for d in daterange(first_day, last_day)])
    (holdout_dir / "split.json").write_text(json.dumps(split, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(holdout_dir / "order_mapping.csv", ["order_name", "product_id", "product_name", "method"],
              sorted([name, ids[target] if target else "", target or "", method] for name, (target, method) in order_target.items()))
    write_csv(holdout_dir / "product_aliases.csv", ["pos_name", "product_id", "product_name"],
              sorted([raw, ids[name], name] for raw, name in canonical.items() if raw != name))
    write_csv(holdout_dir / "product_stock_check.csv", ["product_id", "name", "category", "sold_units", "ordered_units", "coverage", "price", "unit_cost", "stock_estimate"],
              [[pid, names[pid], categories[pid], sum(sold[pid].values()), order_units[pid], round(profile[pid]["coverage"], 3), profile[pid]["price"], profile[pid]["cost"] or "", "사용" if profile[pid]["reliable"] else ""] for pid in sorted(names)])

    window = daterange(holdout_start, demo_end)
    # 이 기간에 팔린 상품 전부. 판매 행은 상품의 첫·마지막 판매일 사이만 둔다.
    prices = {pid: profile[pid]["price"] or mode_price(price_rows[pid], last_day) for pid in names}
    chosen = sorted((pid for pid in names if prices[pid] and any(sold[pid].get(d, 0) > 0 for d in window)),
                    key=lambda pid: -sum(sold[pid][d] for d in window))
    span = {pid: (min(d for d in sold[pid] if sold[pid][d] > 0), max(d for d in sold[pid] if sold[pid][d] > 0)) for pid in chosen}
    sales_rows = sum(span[pid][0] <= d <= span[pid][1] for pid in chosen for d in window)
    if len(chosen) > 400 or sales_rows > 30_000:
        raise ValueError("시연 자료가 업로드 한도(상품 400개, 30,000행)를 넘습니다.")
    # 발주 기록으로 알 수 없으면 원가는 판매가의 70%, 재고는 최근 1주 평균 판매의 3일치로 가정한다.
    costs = {pid: profile[pid]["cost"] if profile[pid]["cost"] and COST_RATIO[0] <= profile[pid]["cost"] / prices[pid] <= COST_RATIO[1]
             else round(prices[pid] * 0.7 / 10) * 10 for pid in chosen}
    assumed_cost = sum(costs[pid] != profile[pid]["cost"] for pid in chosen)
    assumed_stock = [pid for pid in chosen if pid not in stock]

    book = Workbook()
    guide = book.active
    guide.title = "안내"
    guide.append(["항목", "설명"])
    guide.append(["매장명 예시", "디미고 매점"])
    guide.append(["지역", "한국디지털미디어고등학교 (경기 안산시 단원구) / Asia/Seoul"])
    guide.append(["마감 기준일", demo_date.isoformat()])
    guide.append(["자료 출처", f"디미고사회적협동조합 POS 판매현황 {first_day}~{last_day} 중 뒤쪽 30% 기간 {holdout_start}~{demo_end}. 요일을 유지하도록 날짜를 {shift.days}일 뒤로 옮겼습니다. 모델은 이 기간을 학습하지 않았습니다."])
    guide.append(["상품", f"이 기간에 팔린 상품 전부({len(chosen)}개). 판매는 상품의 첫·마지막 판매일 사이만 담았습니다."])
    guide.append(["재고", f"실사 재고가 아닙니다. 발주 기록이 판매와 맞는 {len(chosen) - len(assumed_stock)}개는 구매현황 누적에서 판매 누적을 빼고 선입선출·종류별 유통기한으로 추정했고, 나머지 {len(assumed_stock)}개는 최근 1주 평균 판매의 3일치로 가정했습니다."])
    guide.append(["가정값", f"원가는 발주 단가(합계÷수량)이며, 발주 기록으로 알 수 없는 {assumed_cost}개는 판매가의 70%로 가정했습니다. 납기 {LEAD_DAYS}일·발주 주기 {REVIEW_DAYS}일·유통기한(종류별)은 시연 가정입니다. 거래처는 가상 유통사 10곳이며, 같은 제조사 상품은 같은 거래처로 묶었습니다."])
    guide.append(["날씨", f"date는 업로드한 날짜, sold_on은 실제로 판매한 날짜({shift.days}일 전)입니다. 기온(°C)·강수(mm)는 실제 판매일의 안산 와동 과거 재분석(Open-Meteo)이고, units_sold_total은 그날 이 {len(chosen)}개 상품 판매 합계입니다. 앱은 매장 위치의 날씨를 date 기준으로 따로 조회하므로 앱의 과거 날씨와 다를 수 있습니다. 업로드에는 쓰지 않는 참고 시트입니다."])
    products_sheet = book.create_sheet("상품")
    products_sheet.append(["id", "name", "category", "supplier", "price", "cost", "lead_days", "shelf_days", "review_days"])
    stocks = book.create_sheet("재고")
    stocks.append(["product_id", "measured_on", "quantity", "expires_on", "arrives_on"])
    sales_sheet = book.create_sheet("판매")
    sales_sheet.append(["date", "product_id", "units_sold"])
    suppliers = assign_distributors({pid: maker_group(mapped_orders[pid], categories[pid]) for pid in chosen})
    for pid in chosen:
        shelf = SHELF_DAYS[categories[pid]]
        products_sheet.append([pid, names[pid], categories[pid], suppliers[pid], prices[pid], costs[pid], LEAD_DAYS, shelf, REVIEW_DAYS])
        batches = Counter()
        if pid in stock:
            for received, quantity in stock[pid][1][demo_end] or [(demo_end, 0)]:
                batches[received + timedelta(days=shelf) + shift] += quantity
        else:
            recent = [sold[pid].get(demo_end - timedelta(days=i), 0) for i in range(7)]
            batches[demo_date + timedelta(days=shelf)] = -(-sum(recent) * 3 // 7)
        for expires_on, quantity in sorted(batches.items()):
            stocks.append([pid, demo_date.isoformat(), quantity, expires_on.isoformat(), None])
        for day in window:
            if span[pid][0] <= day <= span[pid][1]:
                sales_sheet.append([(day + shift).isoformat(), pid, sold[pid].get(day, 0)])
    # 날짜를 옮겨도 판매와 날씨의 짝은 실제 판매일 기준으로 남긴다. 업로드에는 쓰지 않는 참고 시트다.
    weather_sheet = book.create_sheet("날씨")
    weather_sheet.append(["date", "sold_on", "weekday", "temp_max", "temp_min", "rain_mm", "units_sold_total"])
    for day in window:
        weather_sheet.append([(day + shift).isoformat(), day.isoformat(), "월화수목금토일"[day.weekday()], *weather[day][:3], sum(sold[pid].get(day, 0) for pid in chosen)])
    styled(book)
    guide.column_dimensions["B"].width = 100
    path = demo_dir / f"디미고매점_매장자료_{demo_date}.xlsx"
    book.save(path)
    (demo_dir / "사용순서.txt").write_text(f"""화개장날 · 디미고 매점 시연 순서

1. 점주 계정으로 로그인 → 매장 선택 옆 + → 매장명 '디미고 매점' → 지역 검색 '한국디지털미디어고등학교' → 매장 추가.
2. '우리 매장 자료를 가져오세요'에서 {path.name} 선택 → 내용 확인하기.
   상품 {len(chosen)}개, 판매 {sales_rows}건, 재고 기준일 {demo_date}를 확인합니다.
   '날씨' 시트에는 날짜별 실제 판매일의 기온·강수와 판매 합계가 있습니다(발표 근거용, 업로드에는 쓰이지 않음).
3. 현재 재고·입고 예정 확인에 체크 → 확인한 자료 저장.
4. 7일 판매 예측 → 실제 안산 예보로 다음 7일 확인. 과거 날짜를 고르면 실제 판매와 비교됩니다.
5. 발주 계획 → 발주량 계산 → 수량 수정·저장 → 발주 확정 → Excel/CSV.

이 파일은 시연일 {demo_date} 전용입니다. 다른 날 시연하면 다시 생성하세요.
uv run python scripts/make-dimigo-data.py --orders <구매 XLSX> --sales <판매 XLSX> --output <폴더> --demo-date <시연일>
""", encoding="utf-8-sig")

    reliable = [pid for pid in names if profile[pid]["reliable"]]
    return {
        "products": len(names), "merged_names": sum(1 for a, b in canonical.items() if a != b),
        "order_names": len(order_target), "mapped_order_names": sum(1 for t, _ in order_target.values() if t),
        "stock_products": len(reliable), "train": split["train"], "holdout": split["holdout"],
        "demo": {"date": demo_date.isoformat(), "source_end": demo_end.isoformat(), "shift_days": shift.days, "products": len(chosen), "sales_rows": sales_rows},
        "output": str(output),
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--orders", type=Path, required=True, help="구매현황집계 XLSX")
    parser.add_argument("--sales", type=Path, required=True, help="판매현황집계 XLSX")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--demo-date", type=date.fromisoformat, default=None, help="시연일 (기본: 오늘, 서울)")
    args = parser.parse_args()
    summary = create_packs(args.orders, args.sales, args.output, args.demo_date or datetime.now(ZoneInfo(TIMEZONE)).date())
    print(json.dumps(summary, ensure_ascii=False, indent=2))
