"""A merchant uploads their own file; preview is read-only and import is all-or-nothing."""

import io
import runpy
from datetime import date, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import store_today
from app.main import create_app
from app.setup_import import read_workbook
from test_walmart_live import Runtime, weather_snapshot

MAKE = runpy.run_path(str(Path(__file__).resolve().parents[3] / "scripts/make-merchant-data.py"))["create_pack"]


@pytest.fixture
def session(tmp_path):
    app = create_app(tmp_path / "imports.sqlite3")
    runtime = Runtime()
    app.state.model_runtime = runtime
    app.state.merchant_model_runtime = Runtime()
    app.state.weather_reader = lambda db, store, cutoff: weather_snapshot(store, cutoff)
    with TestClient(app) as client:
        auth = {}
        for role in ("owner", "staff"):
            token = client.post("/auth/login", json={"email": f"{role}@hwagae.local", "password": "hwagae-demo"}).json()["token"]
            auth[role] = {"Authorization": "Bearer " + token}
        store_id = client.post("/stores", headers=auth["owner"], json={"name": "화개마켓", "city": "댈러스", "latitude": 32.7767, "longitude": -96.797, "timezone": "America/Chicago"}).json()["id"]
        cutoff = store_today({"timezone": "America/Chicago"})
        content = MAKE(tmp_path / "pack", cutoff).read_bytes()
        yield client, auth, "/stores/" + store_id, content, app, cutoff


def post(client, path, auth, content, *, save=False, confirmed=True):
    return client.post(path + ("/initial-data" if save else "/initial-data/preview"), headers=auth,
                       files={"file": ("매장자료.xlsx", content)}, data={"confirmed": str(confirmed).lower()})


def edit(content, sheet, cell, value):
    book = load_workbook(io.BytesIO(content))
    book[sheet][cell] = value
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_upload_to_forecast_order_and_export(session):
    client, auth, path, content, app, cutoff = session
    preview = post(client, path, auth["owner"], content)
    assert preview.status_code == 200, preview.text
    assert preview.json()["product_count"] == 6
    assert preview.json()["sales_count"] == 672
    assert preview.json()["inventory"]["walmart-4-9"]["batches"][0]["quantity"] == 48
    assert client.get(path + "/products", headers=auth["owner"]).json()["products"] == []
    imported = post(client, path, auth["owner"], content, save=True)
    assert imported.status_code == 201, imported.text
    settings = client.get(path + "/forecast-settings", headers=auth["owner"]).json()
    assert settings["model_kind"] == "walmart" and settings["provider"] == "merchant_tft"
    result = client.post(path + "/forecasts", headers=auth["owner"])
    assert result.status_code == 200, result.text
    forecast = result.json()
    assert len(app.state.model_runtime.calls) == 1
    assert all(row["status"] == "ready" for row in forecast["products"])
    rec = client.post(path + "/recommendations", headers=auth["owner"], json={"forecast_id": forecast["id"], "order_date": (cutoff + timedelta(days=2)).isoformat()})
    assert rec.status_code == 200 and rec.json()["calculated_count"] == 6
    saved = client.post(path + "/orders", headers=auth["owner"], json={"recommendation_id": rec.json()["id"], "lines": [{"product_id": r["product_id"], "quantity": r["quantity"] + 1} for r in rec.json()["rows"]]})
    assert saved.status_code == 201
    order = saved.json()
    assert client.post(path + "/orders/" + order["id"] + "/confirm", headers=auth["owner"]).status_code == 200
    csv = client.get(path + "/orders/" + order["id"] + "/export?format=csv", headers=auth["owner"])
    assert "햄치즈 샌드위치" in csv.content.decode("utf-8-sig")
    assert post(client, path, auth["owner"], content, save=True).status_code == 400
    assert len(client.get(path + "/products", headers=auth["owner"]).json()["products"]) == 6


@pytest.mark.parametrize("sheet,cell,value", [
    ("상품", "E2", -1), ("상품", "B2", "=1+1"),
    ("판매", "C2", 999), ("판매", "B2", "unknown"),
    ("재고", "B2", "2020-01-01"), ("재고", "C2", -1),
    ("재고", "E3", "2020-01-01"), ("모델연결", "A2", "unknown"),
])
def test_invalid_file_never_partially_imports(session, sheet, cell, value):
    client, auth, path, content, _, _ = session
    bad = edit(content, sheet, cell, value)
    assert post(client, path, auth["owner"], bad, save=True).status_code == 400
    assert client.get(path + "/products", headers=auth["owner"]).json()["products"] == []


def test_owner_confirmation_and_date_constraints(session):
    client, auth, path, content, _, cutoff = session
    assert post(client, path, auth["staff"], content, save=True).status_code == 403
    assert post(client, path, auth["owner"], content, save=True, confirmed=False).status_code == 422
    assert post(client, path, auth["owner"], b"not a zip").status_code == 400
    future = edit(content, "판매", "A2", (cutoff + timedelta(days=1)).isoformat())
    assert post(client, path, auth["owner"], future).status_code == 400


def test_regular_data_uses_general_tft_without_walmart_id_mapping(session):
    client, auth, path, content, _, _ = session
    book = load_workbook(io.BytesIO(content))
    del book["모델연결"]
    output = io.BytesIO(); book.save(output); book.close()
    assert post(client, path, auth["owner"], output.getvalue(), save=True).status_code == 201
    settings = client.get(path + "/forecast-settings", headers=auth["owner"]).json()
    assert settings["model_kind"] == "dimigo_store"


def test_recheck_empty_store_inside_transaction(session):
    from app.setup_import import persist
    client, auth, path, content, app, cutoff = session
    store = {"id": path.split("/")[-1], "timezone": "America/Chicago"}
    data = read_workbook("data.xlsx", content, cutoff)
    data["cutoff"] = cutoff.isoformat()
    assert post(client, path, auth["owner"], content, save=True).status_code == 201
    with pytest.raises(ValueError, match="새 매장"):
        persist(app.state.database, store, data)


def test_date_change_between_preview_and_commit_is_rejected(session, monkeypatch):
    from app.setup_import import persist
    client, auth, path, content, app, cutoff = session
    data = read_workbook("data.xlsx", content, cutoff)
    data["cutoff"] = cutoff.isoformat()
    monkeypatch.setattr("app.setup_import.store_today", lambda store: cutoff + timedelta(days=1))
    with pytest.raises(ValueError, match="날짜가 바뀌었습니다"):
        persist(app.state.database, {"id": path.split("/")[-1]}, data)
    assert client.get(path + "/products", headers=auth["owner"]).json()["products"] == []
