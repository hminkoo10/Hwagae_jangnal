"""Walmart date-shifted US scenarios remain separate from historical evaluation."""

import csv
import io
from copy import deepcopy
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database import now_iso, store_today
from app.main import create_app
from app.walmart_scenarios import LOCATIONS, scenario_history


class Runtime:
    def __init__(self):
        self.calls = []

    def start(self):
        pass

    def close(self):
        pass

    def status(self):
        return {"status": "ready", "message": "test runtime"}

    def predict(self, context):
        self.calls.append(deepcopy(context))
        return {
            "model": {
                "name": "Walmart TFT · 미국 날씨 시나리오",
                "version": "walmart-test",
                "uses_weather": True,
                "training_data": "walmart",
                "experimental": True,
                "weather_strategy": "provided_forecast",
                "note": "날짜 이동 판매와 선택 지역 날씨 시연. 기존 평가 미적용.",
            },
            "products": [
                {
                    "product_id": product["id"],
                    "status": "ready",
                    "reason": None,
                    "points": [
                        {"date": day, "p10": 20, "p50": 30, "p90": 40}
                        for day in context["dates"]
                    ],
                }
                for product in context["products"]
            ],
        }


def weather_snapshot(store, cutoff):
    return {
        "status": "ok",
        "source": "Open-Meteo · NOAA GFS/HRRR fixture",
        "mode": "live_forecast" if cutoff == store_today(store) else "historical_reanalysis",
        "timezone": store["timezone"],
        "units": {"temperature": "celsius", "precipitation": "mm"},
        "fetched_at": now_iso(),
        "cached": False,
        "days": [
            {
                "date": (cutoff + timedelta(days=i)).isoformat(),
                "temp_max": 30 + i - 1,
                "temp_min": 10 + i - 1,
                "temp_mean": 20 + i - 1,
                "rain_mm": 25.4 * i,
                "wind_mean_ms": 4.4704 * i,
            }
            for i in range(1, 8)
        ],
    }


@pytest.fixture
def app_client(tmp_path):
    app = create_app(tmp_path / "walmart-us.sqlite3")
    walmart, merchant = Runtime(), Runtime()
    app.state.model_runtime = walmart
    app.state.merchant_model_runtime = merchant
    app.state.weather_reader = lambda db, store, cutoff: weather_snapshot(store, cutoff)
    with TestClient(app) as client:
        auth = {}
        for role in ("owner", "staff"):
            response = client.post(
                "/auth/login",
                json={"email": f"{role}@hwagae.local", "password": "hwagae-demo"},
            )
            assert response.status_code == 200
            auth[role] = {"Authorization": "Bearer " + response.json()["token"]}
        yield client, auth, walmart, merchant, app.state.database


def create_store(client, auth, location="new_york", sample_id="walmart-demo-4"):
    response = client.post(
        "/stores/test-walmart",
        headers=auth,
        json={"name": "Walmart US test", "location": location, "sample_id": sample_id},
    )
    assert response.status_code == 201, response.text
    return response.json()


@pytest.mark.parametrize(
    "location,sample_id",
    [
        ("dallas", "walmart-demo-4"),
        ("new_york", "walmart-demo-24"),
        ("seattle", "walmart-demo-17"),
    ],
)
def test_walmart_store_preserves_existing_data_and_marks_date_shift(
    app_client, location, sample_id
):
    client, auth, _, _, database = app_client
    tables = ("stores", "products", "sales", "inventory")
    with database.connect() as db:
        before = {
            table: {tuple(row) for row in db.execute(f"SELECT * FROM {table}")}
            for table in tables
        }
    store = create_store(client, auth["owner"], location, sample_id)
    assert store["timezone"] == LOCATIONS[location]["timezone"]
    cutoff = store_today(LOCATIONS[location])
    assert store["local_today"] == cutoff.isoformat()
    assert "날짜" in store["provenance"] and "실제 위치가 아닙니다" in store["provenance"]
    source = next(item for item in scenario_history()["stores"] if item["id"] == sample_id)
    with database.connect() as db:
        for table in tables:
            assert before[table] <= {tuple(row) for row in db.execute(f"SELECT * FROM {table}")}
        count, days, first, last, provenance = db.execute(
            "SELECT COUNT(*),COUNT(DISTINCT date),MIN(date),MAX(date),MIN(source) "
            "FROM sales WHERE store_id=?",
            (store["id"],),
        ).fetchone()
        assert (count, days) == (112 * len(source["products"]), 112)
        assert first == (cutoff - timedelta(days=111)).isoformat()
        assert last == cutoff.isoformat() and provenance == "walmart_date_shifted_demo"
        actual = [
            (row[0], row[1])
            for row in db.execute(
                "SELECT product_id,units_sold FROM sales "
                "WHERE store_id=? ORDER BY date,product_id",
                (store["id"],),
            )
        ]
        expected = [
            (row["product_id"], row["units_sold"])
            for row in sorted(
                source["sales"], key=lambda row: (row["date"], row["product_id"])
            )
        ]
        assert actual == expected
    assert client.get(f"/stores/{store['id']}/products", headers=auth["staff"]).status_code == 403
    settings = client.get(f"/stores/{store['id']}/forecast-settings", headers=auth["owner"]).json()
    assert settings["provider"] == "merchant_tft" and settings["model_kind"] == "walmart"


def test_walmart_create_requires_owner_and_strict_fixed_inputs(app_client):
    client, auth, _, _, database = app_client
    body = {"name": "test", "location": "dallas", "sample_id": "walmart-demo-4"}
    assert client.post("/stores/test-walmart", json=body).status_code == 401
    assert client.post("/stores/test-walmart", headers=auth["staff"], json=body).status_code == 403
    for changes in (
        {"location": "unknown"},
        {"sample_id": "../input.json"},
        {"latitude": 50},
        {"name": ""},
    ):
        assert client.post(
            "/stores/test-walmart", headers=auth["owner"], json={**body, **changes}
        ).status_code == 422
    with database.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM walmart_scenarios").fetchone()[0] == 0


def test_forecast_uses_walmart_runtime_converted_us_weather_and_observed_history(app_client):
    client, auth, walmart, merchant, _ = app_client
    store = create_store(client, auth["owner"])
    response = client.post(f"/stores/{store['id']}/forecasts", headers=auth["owner"])
    assert response.status_code == 200, response.text
    result = response.json()
    assert merchant.calls == [] and len(walmart.calls) == 1
    context = walmart.calls[0]
    assert context["as_of"] == store["local_today"]
    assert all(row["date"] <= context["as_of"] for row in context["sales"])
    assert all(
        point["actual"] is None
        for product in result["products"]
        for point in product["points"]
    )
    spec = context["model_input"]
    assert spec["weather_strategy"] == "provided_forecast" and spec["weather_units"] == "walmart"
    assert all(row["date"] <= context["as_of"] for row in spec["history_weather"])
    day = spec["forecast_weather"][0]
    assert day["tavg"] == pytest.approx(68)
    assert day["tmax"] == pytest.approx(86) and day["tmin"] == pytest.approx(50)
    assert day["preciptotal"] == pytest.approx(1) and day["avgspeed"] == pytest.approx(10)
    assert [row["date"] for row in spec["forecast_weather"]] == context["dates"]
    assert [row["preciptotal"] for row in spec["forecast_weather"]] == pytest.approx(
        list(range(1, 8))
    )
    assert [row["avgspeed"] for row in spec["forecast_weather"]] == pytest.approx(
        list(range(10, 71, 10))
    )
    source = next(item for item in scenario_history()["stores"] if item["id"] == "walmart-demo-4")
    assert spec["store_nbr"] == source["store_nbr"] and spec["station_nbr"] == source["station_nbr"]
    assert [row["tmax"] for row in spec["history_weather"]] == [
        row["tmax"] for row in source["history_weather"]
    ]
    assert result["model"]["experimental"] and result["model"]["training_data"] == "walmart"


@pytest.mark.parametrize("mode", ["historical_reanalysis", "recent_weather"])
def test_past_scenarios_use_past_weather_without_inventing_future_sales(app_client, mode):
    client, auth, walmart, merchant, _ = app_client
    store = create_store(client, auth["owner"])
    cutoff = date.fromisoformat(store["local_today"]) - timedelta(days=14)

    def reader(db, current, origin):
        return {**weather_snapshot(current, origin), "mode": mode}

    client.app.state.weather_reader = reader
    prefix = f"/stores/{store['id']}"
    response = client.post(
        prefix + "/forecasts",
        params={"as_of": cutoff.isoformat()},
        headers=auth["owner"],
    )
    assert response.status_code == 200, response.text
    assert all(row["date"] <= cutoff.isoformat() for row in walmart.calls[0]["sales"])
    assert merchant.calls == []
    client.app.state.weather_reader = lambda db, current, origin: {
        **weather_snapshot(current, origin), "mode": "live_forecast"
    }
    assert client.post(
        prefix + "/forecasts",
        params={"as_of": cutoff.isoformat()},
        headers=auth["owner"],
    ).status_code == 503
    tomorrow = (date.fromisoformat(store["local_today"]) + timedelta(days=1)).isoformat()
    assert client.post(
        prefix + "/forecasts", params={"as_of": tomorrow}, headers=auth["owner"]
    ).status_code == 422
    assert len(walmart.calls) == 1


def test_future_order_can_be_edited_saved_confirmed_and_exported(app_client):
    client, auth, _, _, _ = app_client
    store = create_store(client, auth["owner"])
    prefix = f"/stores/{store['id']}"
    forecast = client.post(prefix + "/forecasts", headers=auth["owner"])
    assert forecast.status_code == 200, forecast.text
    cutoff = date.fromisoformat(store["local_today"])
    order_date = (cutoff + timedelta(days=2)).isoformat()
    response = client.post(
        prefix + "/recommendations",
        headers=auth["owner"],
        json={"forecast_id": forecast.json()["id"], "order_date": order_date},
    )
    assert response.status_code == 200, response.text
    recommendation = response.json()
    assert recommendation["order_date"] == order_date
    assert all(
        row["arrival_date"] == (cutoff + timedelta(days=3)).isoformat()
        for row in recommendation["rows"]
    )
    ready = next(row for row in recommendation["rows"] if row["status"] == "ready")
    edited_quantity = ready["quantity"] + 3
    response = client.post(
        prefix + "/orders",
        headers=auth["owner"],
        json={
            "recommendation_id": recommendation["id"],
            "lines": [
                {"product_id": ready["product_id"], "quantity": edited_quantity}
            ],
        },
    )
    assert response.status_code == 201, response.text
    order = response.json()
    assert order["order_date"] == order_date and order["lines"][0]["quantity"] == edited_quantity
    url = prefix + f"/orders/{order['id']}"
    assert client.post(url + "/confirm", headers=auth["owner"]).status_code == 200
    csv_response = client.get(url + "/export?format=csv", headers=auth["owner"])
    rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    column = rows[0].index("발주예정일") if "발주예정일" in rows[0] else rows[0].index("발주 예정일")
    assert rows[1][column] == order_date
    sheet = load_workbook(
        io.BytesIO(client.get(url + "/export?format=xlsx", headers=auth["owner"]).content)
    ).active
    assert sheet.cell(2, column + 1).value == order_date
    # Latest forecast covers seven days; an order on D+6 with a two-day review
    # window needs D+8 and cannot silently assume demand beyond the horizon.
    response = client.post(
        prefix + "/recommendations",
        headers=auth["owner"],
        json={
            "forecast_id": forecast.json()["id"],
            "order_date": (cutoff + timedelta(days=6)).isoformat(),
        },
    )
    assert response.status_code == 200
    assert all(
        row["status"] == "blocked" and "7일" in row["reason"]
        for row in response.json()["rows"]
    )
