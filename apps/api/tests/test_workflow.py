"""Store isolation, data integrity, model handoff and complete ordering workflow."""

import csv
import io
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.database import today
from app.main import create_app
from app.ordering import simulate


@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path / "test.sqlite3")
    app.state.weather_reader = lambda db, store, cutoff: {
        "status": "ok",
        "source": "test_fixture",
        "mode": "live_forecast" if cutoff == today() else "historical_reanalysis",
        "days": [],
        "fetched_at": "2026-10-24T00:00:00Z",
        "note": "test",
    }
    with TestClient(app) as session:
        yield session


def headers(client, role="owner"):
    response = client.post(
        "/auth/login", json={"email": f"{role}@hwagae.local", "password": "hwagae-demo"}
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['token']}"}


def forecast(client, auth, store="gwangju", as_of=None):
    response = client.post(
        f"/stores/{store}/forecasts",
        headers=auth,
        params={"as_of": as_of} if as_of else {},
    )
    assert response.status_code == 200, response.text
    return response.json()


def stock(client, auth, quantity=0):
    response = client.put(
        "/stores/gwangju/inventory/P001",
        headers=auth,
        json={
            "batches": [
                {
                    "quantity": quantity,
                    "expires_on": (today() + timedelta(days=2)).isoformat(),
                }
            ],
            "incoming": [],
            "incoming_confirmed": True,
        },
    )
    assert response.status_code == 200, response.text


def recommendation(client, auth, prediction):
    response = client.post(
        "/stores/gwangju/recommendations",
        headers=auth,
        json={"forecast_id": prediction["id"]},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_authentication_and_store_isolation(client):
    assert client.get("/stores").status_code == 401
    auth = headers(client, "staff")
    assert [
        store["id"] for store in client.get("/stores", headers=auth).json()["stores"]
    ] == ["gwangju"]
    assert client.get("/stores/busan/products", headers=auth).status_code == 403
    assert client.post("/stores/busan/forecasts", headers=auth).status_code == 403
    assert client.post("/auth/logout", headers=auth).status_code == 200
    assert client.get("/stores", headers=auth).status_code == 401


def test_null_inventory_is_not_zero_and_zero_unlocks_ordering(client):
    auth = headers(client)
    prediction = forecast(client, auth)
    assert prediction["model"]["uses_weather"] is False
    assert all(len(product["points"]) == 7 for product in prediction["products"])
    initial = recommendation(client, auth, prediction)
    assert initial["calculated_count"] == 0
    assert all(row["quantity"] is None for row in initial["rows"])
    stock(client, auth, 0)
    updated = recommendation(client, auth, prediction)
    first = next(row for row in updated["rows"] if row["product_id"] == "P001")
    assert updated["calculated_count"] == 1
    assert first["quantity"] > 0
    assert first["current_stock"] == 0


def test_csv_preview_blocks_errors_without_writing(client):
    auth = headers(client)
    raw = "date,product_id,units_sold\n2020-01-01,P001,5\n2020-01-02,NOPE,3\n2020-01-03,P001,-1\n2020-01-01,P001,7\n".encode()
    response = client.post(
        "/stores/gwangju/sales/preview",
        headers=auth,
        files={"file": ("sales.csv", raw, "text/csv")},
    )
    assert response.status_code == 200
    assert not response.json()["can_save"]
    assert len(response.json()["errors"]) == 3
    assert (
        client.get("/stores/gwangju/sales", headers=auth).json()["summary"]["count"]
        == 840
    )


def test_csv_and_xlsx_import_preserve_integer_zero(client):
    auth = headers(client)
    for filetype, day in (("csv", "2020-01-01"), ("xlsx", "2020-01-02")):
        if filetype == "csv":
            content = f"date,product_id,units_sold\n{day},P001,0\n".encode("utf-8-sig")
        else:
            book = Workbook()
            book.active.append(["날짜", "상품코드", "판매수량"])
            book.active.append([day, "P001", 0])
            stream = io.BytesIO()
            book.save(stream)
            content = stream.getvalue()
        result = client.post(
            "/stores/gwangju/sales/preview",
            headers=auth,
            files={"file": (f"sales.{filetype}", content)},
        )
        assert result.status_code == 200, result.text
        assert result.json()["can_save"]
        rows = result.json()["rows"]
        assert rows[0]["units_sold"] == 0
        assert (
            client.post(
                "/stores/gwangju/sales", headers=auth, json={"rows": rows}
            ).status_code
            == 201
        )


def test_duplicate_commit_rolls_back_entire_batch(client):
    auth = headers(client)
    existing = client.get("/stores/gwangju/sales?limit=1", headers=auth).json()["rows"][
        0
    ]
    response = client.post(
        "/stores/gwangju/sales",
        headers=auth,
        json={
            "rows": [
                {"date": "2020-02-01", "product_id": "P001", "units_sold": 8},
                {key: existing[key] for key in ("date", "product_id", "units_sold")},
            ]
        },
    )
    assert response.status_code == 409
    assert (
        client.get("/stores/gwangju/sales", headers=auth).json()["summary"]["count"]
        == 840
    )


def test_replay_hides_future_from_inputs_and_cannot_order(client):
    auth = headers(client)
    old = (today() - timedelta(days=14)).isoformat()
    before = forecast(client, auth, as_of=old)
    assert before["products"][0]["points"][0]["actual"] is not None
    # Changing a later outcome must not change the earlier prediction.
    db = client.app.state.database
    with db.connect() as connection:
        connection.execute(
            "UPDATE sales SET units_sold=99999 WHERE store_id='gwangju' AND date>?",
            (old,),
        )
    after = forecast(client, auth, as_of=old)
    assert [point["p50"] for point in before["products"][0]["points"]] == [
        point["p50"] for point in after["products"][0]["points"]
    ]
    assert (
        client.post(
            "/stores/gwangju/recommendations",
            headers=auth,
            json={"forecast_id": before["id"]},
        ).status_code
        == 409
    )


def test_save_confirm_export_and_stale_plan_protection(client):
    auth = headers(client)
    stock(client, auth)
    rec = recommendation(client, auth, forecast(client, auth))
    created = client.post(
        "/stores/gwangju/orders",
        headers=auth,
        json={
            "recommendation_id": rec["id"],
            "lines": [{"product_id": "P001", "quantity": 17}],
        },
    )
    assert created.status_code == 201, created.text
    order = created.json()
    assert order["total"] == 17 * 2600
    path = f"/stores/gwangju/orders/{order['id']}"
    assert client.get(path + "/export", headers=auth).status_code == 409
    staff = headers(client, "staff")
    assert client.post(path + "/confirm", headers=staff).status_code == 403
    assert client.post(path + "/confirm", headers=auth).status_code == 200
    csv_result = client.get(path + "/export?format=csv", headers=auth)
    assert "17" in csv_result.content.decode("utf-8-sig")
    workbook = load_workbook(
        io.BytesIO(client.get(path + "/export?format=xlsx", headers=auth).content)
    )
    assert workbook.active["G2"].value == 17
    assert (
        client.get("/stores/gwangju/orders", headers=auth).json()["orders"][0]["lines"][
            0
        ]["quantity"]
        == 17
    )
    stock(client, auth, 100)
    assert (
        client.post(
            "/stores/gwangju/orders",
            headers=auth,
            json={
                "recommendation_id": rec["id"],
                "lines": [{"product_id": "P001", "quantity": 1}],
            },
        ).status_code
        == 409
    )


@pytest.mark.parametrize("cause", ["date", "inventory"])
def test_confirmation_explains_outdated_draft_without_confirming(client, monkeypatch, cause):
    auth = headers(client)
    stock(client, auth)
    rec = recommendation(client, auth, forecast(client, auth))
    created = client.post(
        "/stores/gwangju/orders", headers=auth,
        json={"recommendation_id": rec["id"], "lines": [{"product_id": "P001", "quantity": 17}]},
    )
    assert created.status_code == 201
    if cause == "date":
        monkeypatch.setattr("app.main.store_today", lambda store: today() + timedelta(days=1))
        message = "오늘 자료로 다시 계산"
    else:
        stock(client, auth, 100)
        message = "정보가 바뀌었어요"
    response = client.post(f"/stores/gwangju/orders/{created.json()['id']}/confirm", headers=auth)
    assert response.status_code == 409
    assert message in response.json()["detail"]
    saved = client.get("/stores/gwangju/orders", headers=auth).json()["orders"][0]
    assert saved["status"] == "draft" and saved["confirmed_at"] is None


def test_expired_stock_and_incoming_are_handled_by_date():
    cutoff = today()
    product = {"lead_days": 2, "shelf_days": 1, "cost": 5, "price": 10}
    inventory = {
        "batches": [{"quantity": 100, "expires_on": cutoff.isoformat()}],
        "incoming": [],
    }
    result = simulate(10, [5, 5], product, inventory, cutoff)
    assert result == {
        "waste_units": 5,
        "missed_units": 5,
        "ending_units": 0,
        "loss": 50,
    }
    incoming = {
        "batches": [],
        "incoming": [
            {
                "quantity": 5,
                "arrives_on": (cutoff + timedelta(days=1)).isoformat(),
                "expires_on": (cutoff + timedelta(days=3)).isoformat(),
            }
        ],
    }
    result = simulate(5, [5, 5], product, incoming, cutoff)
    assert result["missed_units"] == result["waste_units"] == 0


def test_empty_store_and_server_validation(client):
    auth = headers(client)
    new = client.post(
        "/stores",
        headers=auth,
        json={"name": "새 매장", "city": "서울", "latitude": 37.5, "longitude": 127.0},
    )
    assert new.status_code == 201
    store_id = new.json()["id"]
    product = {
        "id": "NEW1",
        "name": "새 상품",
        "category": "식품",
        "supplier": "공급사",
        "price": 3000,
        "cost": 1500,
        "lead_days": 1,
        "shelf_days": 3,
        "review_days": 1,
    }
    assert (
        client.post(
            f"/stores/{store_id}/products", headers=auth, json=product
        ).status_code
        == 201
    )
    assert (
        forecast(client, auth, store_id)["products"][0]["status"]
        == "insufficient_history"
    )
    assert (
        client.post(
            f"/stores/{store_id}/sales",
            headers=auth,
            json={
                "rows": [
                    {
                        "date": today().isoformat(),
                        "product_id": "NEW1",
                        "units_sold": -1,
                    }
                ]
            },
        ).status_code
        == 422
    )

    assert (
        client.post(
            f"/stores/{store_id}/sales",
            headers=auth,
            json={
                "rows": [
                    {
                        "date": (today() + timedelta(days=1)).isoformat(),
                        "product_id": "NEW1",
                        "units_sold": 1,
                    }
                ]
            },
        ).status_code
        == 422
    )


@pytest.mark.parametrize("filetype", ["csv", "xlsx"])
def test_export_cleans_legacy_xml_characters_before_formula_protection(
    client, filetype
):
    auth = headers(client)
    # Legacy values must stay exportable, even ones that become formulas after XML cleanup.
    with client.app.state.database.connect() as db:
        db.execute("UPDATE stores SET name=? WHERE id='gwangju'", ("\x00=1+1",))
        db.execute(
            "UPDATE products SET name=?,supplier=? WHERE store_id='gwangju' AND id='P001'",
            ("\ufffe+상품", "공급\x08사\uffff\t정상 🍞"),
        )
    stock(client, auth)
    rec = recommendation(client, auth, forecast(client, auth))
    created = client.post(
        "/stores/gwangju/orders",
        headers=auth,
        json={
            "recommendation_id": rec["id"],
            "lines": [{"product_id": "P001", "quantity": 17}],
        },
    )
    assert created.status_code == 201, created.text
    path = f"/stores/gwangju/orders/{created.json()['id']}"
    assert client.post(path + "/confirm", headers=auth).status_code == 200
    result = client.get(path + f"/export?format={filetype}", headers=auth)
    assert result.status_code == 200, result.text
    if filetype == "csv":
        values = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))[1]
    else:
        book = load_workbook(io.BytesIO(result.content), data_only=False)
        try:
            cells = list(book.active[2])
            assert cells[0].data_type == cells[5].data_type == "s"
            values = [cell.value for cell in cells]
        finally:
            book.close()
    assert values[0] == "'=1+1"
    assert values[2] == "공급사\t정상 🍞"
    assert values[5] == "'+상품"
    assert int(values[6]) == 17
    assert int(values[8]) == 17 * 2600


def test_seven_day_maximum_demand_recommendation_can_be_saved(client):
    auth = headers(client)
    product = {
        "id": "HIGH",
        "name": "큰 수량 상품",
        "category": "식품",
        "supplier": "공급사",
        "price": 100,
        "cost": 50,
        "lead_days": 1,
        "shelf_days": 7,
        "review_days": 7,
    }
    assert (
        client.post("/stores/gwangju/products", headers=auth, json=product).status_code
        == 201
    )
    rows = [
        {
            "product_id": "HIGH",
            "date": (today() - timedelta(days=offset)).isoformat(),
            "units_sold": 1_000_000,
        }
        for offset in range(1, 29)
    ]
    assert (
        client.post(
            "/stores/gwangju/sales", headers=auth, json={"rows": rows}
        ).status_code
        == 201
    )
    assert (
        client.put(
            "/stores/gwangju/inventory/HIGH",
            headers=auth,
            json={"batches": [], "incoming": [], "incoming_confirmed": True},
        ).status_code
        == 200
    )
    rec = recommendation(client, auth, forecast(client, auth))
    row = next(row for row in rec["rows"] if row["product_id"] == "HIGH")
    assert row["quantity"] == 7_000_000
    payload = {
        "recommendation_id": rec["id"],
        "lines": [{"product_id": "HIGH", "quantity": row["quantity"]}],
    }
    saved = client.post("/stores/gwangju/orders", headers=auth, json=payload)
    assert saved.status_code == 201, saved.text
    assert saved.json()["total"] == 7_000_000 * 50
    payload["lines"][0]["quantity"] += 1
    assert (
        client.post("/stores/gwangju/orders", headers=auth, json=payload).status_code
        == 422
    )
    # Raising the order horizon limit must not raise the daily sales limit.
    assert (
        client.post(
            "/stores/gwangju/sales",
            headers=auth,
            json={
                "rows": [
                    {
                        "product_id": "HIGH",
                        "date": today().isoformat(),
                        "units_sold": 1_000_001,
                    }
                ]
            },
        ).status_code
        == 422
    )


@pytest.mark.parametrize("change", ["product", "inventory"])
def test_data_changed_during_recommendation_requires_recalculation(
    client, monkeypatch, change
):
    auth = headers(client)
    prediction = forecast(client, auth)
    database = client.app.state.database
    read_stock = database.stock
    entered, resume = threading.Event(), threading.Event()

    def paused_stock(store_id):
        entered.set()
        assert resume.wait(5), "Concurrent product creation did not finish"
        return read_stock(store_id)

    monkeypatch.setattr(database, "stock", paused_stock)
    with ThreadPoolExecutor(max_workers=1) as pool:
        pending = pool.submit(
            client.post,
            "/stores/gwangju/recommendations",
            headers=auth,
            json={"forecast_id": prediction["id"]},
        )
        try:
            assert entered.wait(5), "Recommendation did not reach the stock read"
            if change == "product":
                created = client.post(
                    "/stores/gwangju/products",
                    headers=auth,
                    json={
                        "id": "NEW",
                        "name": "동시에 추가한 상품",
                        "category": "식품",
                        "supplier": "공급사",
                        "price": 100,
                        "cost": 50,
                        "lead_days": 1,
                        "shelf_days": 2,
                        "review_days": 1,
                    },
                )
                assert created.status_code == 201, created.text
            else:
                stock(client, auth, 10)
        finally:
            resume.set()
        result = pending.result(timeout=10)
    assert result.status_code == 409, result.text
    assert "다시" in result.json()["detail"]
    with database.connect() as db:
        assert db.execute("SELECT COUNT(*) FROM recommendations").fetchone()[0] == 0
