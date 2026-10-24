"""Personal historical order lifecycle without torch, weather calls or merchant writes."""

import csv
import io
import json
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from pathlib import Path
from threading import Event

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import model_demo, model_orders
from app.main import create_app
from app.model_runtime import ModelUnavailable

SAMPLE = "walmart-demo-4"
EMPTY = {"batches": [], "incoming": [], "incoming_confirmed": True}
MERCHANT = (
    "stores",
    "products",
    "sales",
    "inventory",
    "forecasts",
    "recommendations",
    "orders",
)


class Runtime:
    def __init__(self):
        root = Path(__file__).resolve().parents[3]
        self.results = {
            f"walmart-demo-{store}": json.loads(
                (root / f"ml/examples/expected-api-store-{store}.json").read_text(
                    encoding="utf-8"
                )
            )
            for store in (4, 24, 17)
        }
        self.calls = []
        self.failure = False

    def start(self):
        pass

    def close(self):
        pass

    def status(self):
        return {"status": "ready", "message": "test"}

    def predict(self, context):
        self.calls.append(deepcopy(context))
        if self.failure:
            raise ModelUnavailable("모델을 사용할 수 없습니다.")
        return deepcopy(self.results[context["store_id"]])


@pytest.fixture
def demo(tmp_path):
    app = create_app(tmp_path / "model-orders.sqlite3")
    runtime = Runtime()
    app.state.model_runtime = runtime
    with TestClient(app) as client:
        auth = {}
        for role in ("owner", "staff"):
            response = client.post(
                "/auth/login",
                json={"email": f"{role}@hwagae.local", "password": "hwagae-demo"},
            )
            assert response.status_code == 200
            auth[role] = {"Authorization": f"Bearer {response.json()['token']}"}
        yield client, auth, runtime, app.state.database


def counts(database, tables=MERCHANT):
    with database.connect() as db:
        return [
            db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
            for table in tables
        ]


def planning(client, auth, sample=SAMPLE):
    response = client.get(f"/model-demo/planning?sample_id={sample}", headers=auth)
    assert response.status_code == 200, response.text
    return response.json()


def inputs(plan):
    return {
        "revision": plan["revision"],
        "items": [
            {
                key: item[key]
                for key in (
                    "product_id",
                    "price",
                    "cost",
                    "lead_days",
                    "shelf_days",
                    "review_days",
                    "stock",
                )
            }
            for item in plan["items"]
        ],
    }


def save(client, auth, payload, sample=SAMPLE):
    response = client.put(
        f"/model-demo/planning?sample_id={sample}", headers=auth, json=payload
    )
    assert response.status_code == 200, response.text
    return response.json()


def forecast(client, auth, sample=SAMPLE):
    response = client.post(f"/model-demo/forecast?sample_id={sample}", headers=auth)
    assert response.status_code == 200, response.text
    return response.json()


def recommendation(client, auth, run, revision):
    response = client.post(
        "/model-demo/recommendations",
        headers=auth,
        json={"forecast_id": run["id"], "revision": revision},
    )
    assert response.status_code == 200, response.text
    return response.json()


def ready(client, auth, sample=SAMPLE):
    payload = inputs(planning(client, auth, sample))
    for item in payload["items"]:
        item["stock"] = deepcopy(EMPTY)
    plan = save(client, auth, payload, sample)
    run = forecast(client, auth, sample)
    rec = recommendation(client, auth, run, plan["revision"])
    return plan, run, rec


def draft(client, auth, rec, quantity=17):
    response = client.post(
        "/model-demo/orders",
        headers=auth,
        json={
            "recommendation_id": rec["id"],
            "lines": [
                {"product_id": rec["rows"][0]["product_id"], "quantity": quantity}
            ],
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_auth_defaults_read_only_and_private_forecast_snapshot(demo):
    client, auth, runtime, database = demo
    for method, url, body in [
        ("get", "/model-demo/planning", None),
        ("put", "/model-demo/planning", {}),
        ("post", "/model-demo/recommendations", {}),
        ("get", "/model-demo/orders", None),
        ("post", "/model-demo/orders", {}),
        ("put", "/model-demo/orders/fake", {}),
        ("post", "/model-demo/orders/fake/confirm", {}),
        ("get", "/model-demo/orders/fake/export", None),
    ]:
        kwargs = {"json": body} if body is not None else {}
        assert getattr(client, method)(url, **kwargs).status_code == 401
    before = counts(database)
    plan = planning(client, auth["owner"])
    assert plan["revision"] == 0 and plan["can_confirm_export"] is True
    assert plan["latest_forecast_id"] is plan["model"] is None
    assert len(plan["items"]) == 6
    assert all(item["stock"] is None for item in plan["items"])
    assert planning(client, auth["staff"])["can_confirm_export"] is False
    assert counts(database, ["model_demo_workspaces"]) == [0]
    assert (
        client.get(
            "/model-demo/planning?sample_id=busan", headers=auth["owner"]
        ).status_code
        == 422
    )
    run = forecast(client, auth["owner"])
    assert all(
        point["actual"] is not None
        for item in run["products"]
        for point in item["points"]
    )
    with database.connect() as db:
        row = db.execute(
            "SELECT * FROM model_demo_runs WHERE id=?", (run["id"],)
        ).fetchone()
        snapshot = json.loads(row["payload"])
        assert row["user_id"] == "demo-owner"
        assert "baseline" not in snapshot
        assert all(
            set(point) == {"date", "p10", "p50", "p90"}
            for item in snapshot["products"]
            for point in item["points"]
        )
    assert planning(client, auth["owner"])["latest_forecast_id"] == run["id"]
    assert planning(client, auth["staff"])["latest_forecast_id"] is None
    assert (
        planning(client, auth["owner"], "walmart-demo-24")["latest_forecast_id"] is None
    )
    for wrong_id in (run["id"], "invented"):
        assert (
            client.post(
                "/model-demo/recommendations",
                headers=auth["staff"],
                json={"forecast_id": wrong_id, "revision": 0},
            ).status_code
            == 404
        )
    assert (
        client.post(
            "/stores/gwangju/recommendations",
            headers=auth["owner"],
            json={"forecast_id": run["id"]},
        ).status_code
        == 404
    )
    assert counts(database) == before
    assert all(
        row["date"] <= runtime.calls[0]["as_of"] for row in runtime.calls[0]["sales"]
    )


def test_unknown_inventory_blocks_but_confirmed_zero_calculates(demo):
    client, auth, _, database = demo
    before = counts(database)
    run = forecast(client, auth["owner"])
    rec = recommendation(client, auth["owner"], run, 0)
    assert rec["calculated_count"] == 0
    assert all(row["quantity"] is row["current_stock"] is None for row in rec["rows"])
    payload = inputs(planning(client, auth["owner"]))
    payload["items"][0]["stock"] = deepcopy(EMPTY)
    plan = save(client, auth["owner"], payload)
    assert plan["items"][0]["stock"] == EMPTY
    rec = recommendation(client, auth["owner"], run, 1)
    assert rec["calculated_count"] == 1
    assert rec["rows"][0]["current_stock"] == 0
    assert rec["rows"][0]["quantity"] > 0
    with database.connect() as db:
        settings = json.loads(
            db.execute("SELECT payload FROM model_demo_workspaces").fetchone()[0]
        )
        assert settings[0]["stock"]["measured_on"] == "2014-09-12"
    payload = inputs(plan)
    payload["items"][0]["price"] = payload["items"][0]["cost"]
    save(client, auth["owner"], payload)
    rec = recommendation(client, auth["owner"], run, 2)
    assert rec["rows"][0]["status"] == "blocked"
    assert counts(database) == before


@pytest.mark.parametrize(
    "case",
    [
        "missing",
        "duplicate",
        "foreign",
        "string_price",
        "boolean_quantity",
        "extra",
        "measured_on",
        "expired",
        "arrived",
        "expiry_before_arrival",
        "unconfirmed",
    ],
)
def test_planning_rejects_invalid_inputs_without_saving(demo, case):
    client, auth, _, database = demo
    payload = inputs(planning(client, auth["owner"]))
    item = payload["items"][0]
    item["stock"] = deepcopy(EMPTY)
    if case == "missing":
        payload["items"].pop()
    elif case == "duplicate":
        payload["items"][-1] = deepcopy(item)
    elif case == "foreign":
        item["product_id"] = "P001"
    elif case == "string_price":
        item["price"] = "4000"
    elif case == "boolean_quantity":
        item["stock"]["batches"] = [{"quantity": True, "expires_on": "2014-09-13"}]
    elif case == "extra":
        item["supplier"] = "forged"
    elif case == "measured_on":
        item["stock"]["measured_on"] = "2099-01-01"
    elif case == "expired":
        item["stock"]["batches"] = [{"quantity": 1, "expires_on": "2014-09-12"}]
    elif case == "arrived":
        item["stock"]["incoming"] = [
            {"quantity": 1, "arrives_on": "2014-09-12", "expires_on": "2014-09-13"}
        ]
    elif case == "expiry_before_arrival":
        item["stock"]["incoming"] = [
            {"quantity": 1, "arrives_on": "2014-09-14", "expires_on": "2014-09-13"}
        ]
    elif case == "unconfirmed":
        item["stock"]["incoming_confirmed"] = False
    response = client.put("/model-demo/planning", headers=auth["owner"], json=payload)
    assert response.status_code == 422, response.text
    assert counts(database, ["model_demo_workspaces"]) == [0]


def test_order_lifecycle_reopen_edit_confirm_and_immutable_exports(demo):
    client, auth, _, database = demo
    before = counts(database)
    plan, run, rec = ready(client, auth["owner"])
    order = draft(client, auth["owner"], rec)
    order_id = order["id"]
    url = f"/model-demo/orders/{order_id}"
    assert order["version"] == 0 and order["status"] == "draft"
    assert order["total"] == 34000 and order["stale"] is False
    assert "scenario_loss" not in order["lines"][0]
    assert "scenario_results" not in order["lines"][0]
    assert client.get("/model-demo/orders", headers=auth["owner"]).json()["orders"] == [
        order
    ]
    assert client.get(url + "/export", headers=auth["owner"]).status_code == 409
    update = {
        "version": 0,
        "lines": [{"product_id": order["lines"][0]["product_id"], "quantity": 23}],
    }
    response = client.put(url, headers=auth["owner"], json=update)
    assert response.status_code == 200, response.text
    edited = response.json()
    assert edited["version"] == 1 and edited["total"] == 46000
    assert client.put(url, headers=auth["owner"], json=update).status_code == 409
    assert (
        client.post(
            url + "/confirm", headers=auth["owner"], json={"version": 0}
        ).status_code
        == 409
    )
    confirmed = client.post(
        url + "/confirm", headers=auth["owner"], json={"version": 1}
    )
    assert confirmed.status_code == 200, confirmed.text
    assert (
        confirmed.json()["status"] == "confirmed" and confirmed.json()["version"] == 2
    )
    assert (
        client.put(
            url, headers=auth["owner"], json={**update, "version": 2}
        ).status_code
        == 409
    )
    assert (
        client.post(
            url + "/confirm", headers=auth["owner"], json={"version": 2}
        ).status_code
        == 409
    )
    csv_response = client.get(url + "/export?format=csv", headers=auth["owner"])
    xlsx_response = client.get(url + "/export?format=xlsx", headers=auth["owner"])
    assert csv_response.status_code == xlsx_response.status_code == 200
    csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    sheet = load_workbook(io.BytesIO(xlsx_response.content), data_only=False).active
    assert csv_rows[7] == [str(cell.value) for cell in sheet[8]]
    assert csv_rows[7][5] == "23" and csv_rows[7][7:] == ["2000", "46000"]
    assert csv_rows[1][1] == "2014-09-12" and csv_rows[2][3] == run["model"]["version"]
    assert "시연" in csv_rows[0][0] and "정답" in csv_rows[4][1]
    save(client, auth["owner"], inputs(plan))
    forecast(client, auth["owner"])
    assert (
        client.get("/model-demo/orders", headers=auth["owner"]).json()["orders"][0][
            "stale"
        ]
        is True
    )
    assert (
        client.get(url + "/export", headers=auth["owner"]).content
        == csv_response.content
    )
    assert counts(database) == before


def test_user_sample_isolation_staff_drafts_and_forgery(demo):
    client, auth, _, _ = demo
    _, run, rec = ready(client, auth["owner"])
    assert (
        client.post(
            "/model-demo/recommendations",
            headers=auth["owner"],
            json={"forecast_id": run["id"], "revision": 1, "points": []},
        ).status_code
        == 422
    )
    base = {
        "recommendation_id": rec["id"],
        "lines": [{"product_id": rec["rows"][0]["product_id"], "quantity": 1}],
    }
    assert (
        client.post("/model-demo/orders", headers=auth["staff"], json=base).status_code
        == 404
    )
    for changed in (
        {**base, "cost": 0},
        {**base, "lines": [{"product_id": "P001", "quantity": 1}]},
        {**base, "lines": base["lines"] * 2},
        {**base, "lines": [{**base["lines"][0], "quantity": 7_000_001}]},
    ):
        assert (
            client.post(
                "/model-demo/orders", headers=auth["owner"], json=changed
            ).status_code
            == 422
        )
    _, _, staff_rec = ready(client, auth["staff"])
    order = draft(client, auth["staff"], staff_rec)
    url = f"/model-demo/orders/{order['id']}"
    assert client.get("/model-demo/orders", headers=auth["owner"]).json() == {
        "orders": []
    }
    assert client.get(
        "/model-demo/orders?sample_id=walmart-demo-24", headers=auth["staff"]
    ).json() == {"orders": []}
    for suffix in ("/export",):
        assert client.get(url + suffix, headers=auth["staff"]).status_code == 403
        assert client.get(url + suffix, headers=auth["owner"]).status_code == 404
    assert (
        client.post(
            url + "/confirm", headers=auth["staff"], json={"version": 0}
        ).status_code
        == 403
    )
    assert (
        client.post(
            url + "/confirm", headers=auth["owner"], json={"version": 0}
        ).status_code
        == 404
    )
    assert (
        client.put(
            url, headers=auth["owner"], json={"version": 0, "lines": base["lines"]}
        ).status_code
        == 404
    )


@pytest.mark.parametrize("change", ["settings", "forecast"])
def test_stale_recommendations_and_drafts_cannot_be_saved_edited_or_confirmed(
    demo, change
):
    client, auth, _, _ = demo
    plan, run, rec = ready(client, auth["owner"])
    order = draft(client, auth["owner"], rec)
    if change == "settings":
        save(client, auth["owner"], inputs(plan))
    else:
        forecast(client, auth["owner"])
    assert (
        client.post(
            "/model-demo/recommendations",
            headers=auth["owner"],
            json={"forecast_id": run["id"], "revision": 1},
        ).status_code
        == 409
    )
    lines = [{"product_id": order["lines"][0]["product_id"], "quantity": 20}]
    assert (
        client.post(
            "/model-demo/orders",
            headers=auth["owner"],
            json={"recommendation_id": rec["id"], "lines": lines},
        ).status_code
        == 409
    )
    assert (
        client.put(
            f"/model-demo/orders/{order['id']}",
            headers=auth["owner"],
            json={"version": 0, "lines": lines},
        ).status_code
        == 409
    )
    assert (
        client.post(
            f"/model-demo/orders/{order['id']}/confirm",
            headers=auth["owner"],
            json={"version": 0},
        ).status_code
        == 409
    )
    assert (
        client.get("/model-demo/orders", headers=auth["owner"]).json()["orders"][0][
            "stale"
        ]
        is True
    )


def test_expiry_and_incoming_change_the_recommendation(demo):
    client, auth, runtime, _ = demo
    for product in runtime.results[SAMPLE]["products"]:
        for point in product["points"]:
            point.update(p10=10, p50=10, p90=10)
    run = forecast(client, auth["owner"])
    payload = inputs(planning(client, auth["owner"]))
    item = payload["items"][0]
    item["lead_days"] = 2
    item["stock"] = {
        **deepcopy(EMPTY),
        "batches": [{"quantity": 20, "expires_on": "2014-09-13"}],
    }
    first = save(client, auth["owner"], payload)
    assert (
        recommendation(client, auth["owner"], run, first["revision"])["rows"][0][
            "quantity"
        ]
        == 10
    )
    payload = inputs(first)
    payload["items"][0]["stock"]["batches"][0]["expires_on"] = "2014-09-14"
    second = save(client, auth["owner"], payload)
    assert (
        recommendation(client, auth["owner"], run, second["revision"])["rows"][0][
            "quantity"
        ]
        == 0
    )
    payload = inputs(second)
    payload["items"][0]["stock"] = {
        **deepcopy(EMPTY),
        "incoming": [
            {"quantity": 10, "arrives_on": "2014-09-14", "expires_on": "2014-09-16"}
        ],
    }
    third = save(client, auth["owner"], payload)
    assert (
        recommendation(client, auth["owner"], run, third["revision"])["rows"][0][
            "quantity"
        ]
        == 0
    )


def test_actuals_never_affect_planning_and_failed_runs_do_not_replace_latest(
    demo, monkeypatch, tmp_path
):
    client, auth, runtime, database = demo
    plan, run, rec = ready(client, auth["owner"])
    source = json.loads(model_demo.ACTUALS_PATH.read_text(encoding="utf-8"))
    for row in source["stores"][SAMPLE]:
        row["units_sold"] = 999999
    actuals = tmp_path / "alternate-actuals.json"
    actuals.write_text(json.dumps(source), encoding="utf-8")
    monkeypatch.setattr(model_demo, "ACTUALS_PATH", actuals)
    changed = forecast(client, auth["owner"])
    assert changed["products"][0]["points"][0]["actual"] == 999999
    assert (
        recommendation(client, auth["owner"], changed, plan["revision"])["rows"]
        == rec["rows"]
    )
    assert runtime.calls[0] == runtime.calls[1]
    runtime.failure = True
    before = counts(database, ["model_demo_runs"])
    assert client.post("/model-demo/forecast", headers=auth["owner"]).status_code == 503
    assert counts(database, ["model_demo_runs"]) == before
    assert planning(client, auth["owner"])["latest_forecast_id"] == changed["id"]


@pytest.mark.parametrize("change", ["settings", "forecast"])
def test_recommendation_rechecks_state_after_calculation(demo, monkeypatch, change):
    client, auth, _, database = demo
    plan, run, _ = ready(client, auth["owner"])
    before = counts(database, ["model_demo_recommendations"])
    entered, release = Event(), Event()
    original = model_orders.recommend

    def paused(*args):
        entered.set()
        assert release.wait(10)
        return original(*args)

    monkeypatch.setattr(model_orders, "recommend", paused)
    with ThreadPoolExecutor(max_workers=1) as pool:
        pending = pool.submit(
            client.post,
            "/model-demo/recommendations",
            headers=auth["owner"],
            json={"forecast_id": run["id"], "revision": plan["revision"]},
        )
        try:
            assert entered.wait(10)
            if change == "settings":
                save(client, auth["owner"], inputs(plan))
            else:
                forecast(client, auth["owner"])
        finally:
            release.set()
        assert pending.result().status_code == 409
    assert counts(database, ["model_demo_recommendations"]) == before


def test_revision_and_order_version_prevent_parallel_overwrites(demo):
    client, auth, _, _ = demo
    plan, _, rec = ready(client, auth["owner"])
    order = draft(client, auth["owner"], rec)
    url = f"/model-demo/orders/{order['id']}"
    with ThreadPoolExecutor(max_workers=2) as pool:
        update = pool.submit(
            client.put,
            url,
            headers=auth["owner"],
            json={
                "version": 0,
                "lines": [
                    {"product_id": order["lines"][0]["product_id"], "quantity": 5}
                ],
            },
        )
        confirm = pool.submit(
            client.post, url + "/confirm", headers=auth["owner"], json={"version": 0}
        )
        assert sorted([update.result().status_code, confirm.result().status_code]) == [
            200,
            409,
        ]
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(
            pool.map(
                lambda _: client.put(
                    "/model-demo/planning", headers=auth["owner"], json=inputs(plan)
                ),
                range(2),
            )
        )
    assert sorted(response.status_code for response in responses) == [200, 409]


def test_export_cleans_legacy_xml_before_formula_protection(demo):
    client, auth, _, database = demo
    _, _, rec = ready(client, auth["owner"])
    order = draft(client, auth["owner"], rec)
    with database.connect() as db:
        snapshot = json.loads(
            db.execute(
                "SELECT payload FROM model_demo_orders WHERE id=?", (order["id"],)
            ).fetchone()[0]
        )
        snapshot["lines"][0]["name"] = "\x00=1+1"
        db.execute(
            "UPDATE model_demo_orders SET payload=? WHERE id=?",
            (json.dumps(snapshot), order["id"]),
        )
    assert (
        client.post(
            f"/model-demo/orders/{order['id']}/confirm",
            headers=auth["owner"],
            json={"version": 0},
        ).status_code
        == 200
    )
    url = f"/model-demo/orders/{order['id']}/export"
    csv_rows = list(
        csv.reader(
            io.StringIO(
                client.get(url, headers=auth["owner"]).content.decode("utf-8-sig")
            )
        )
    )
    workbook = load_workbook(
        io.BytesIO(client.get(url + "?format=xlsx", headers=auth["owner"]).content)
    )
    assert csv_rows[7][2] == workbook.active.cell(8, 3).value == "'=1+1"
    assert workbook.active.cell(8, 3).data_type == "s"


def test_confirmation_flag_requires_json_true(demo):
    client, auth, _, _ = demo
    payload = inputs(planning(client, auth["owner"]))
    payload["items"][0]["stock"] = {**deepcopy(EMPTY), "incoming_confirmed": 1}
    response = client.put("/model-demo/planning", headers=auth["owner"], json=payload)
    assert response.status_code == 422


def test_three_sample_workspaces_and_order_reopening_after_app_restart(demo):
    client, auth, _, database = demo
    before = counts(database)
    saved = {}
    for sample, count in [(SAMPLE, 6), ("walmart-demo-24", 3), ("walmart-demo-17", 1)]:
        plan, _, rec = ready(client, auth["owner"], sample)
        assert len(plan["items"]) == rec["calculated_count"] == count
        order = draft(client, auth["owner"], rec)
        assert order["sample_id"] == sample
        saved[sample] = order
    reopened = create_app(database.path)
    reopened.state.model_runtime = Runtime()
    with TestClient(reopened) as second:
        for sample, order in saved.items():
            response = second.get(
                f"/model-demo/orders?sample_id={sample}", headers=auth["owner"]
            )
            assert response.status_code == 200
            assert response.json() == {"orders": [order]}
    assert counts(database) == before


def test_order_edit_preserves_product_set_and_restores_metrics_only_at_recommended_quantity(
    demo,
):
    client, auth, _, _ = demo
    run = forecast(client, auth["owner"])
    blocked = recommendation(client, auth["owner"], run, 0)
    request = {
        "recommendation_id": blocked["id"],
        "lines": [{"product_id": blocked["rows"][0]["product_id"], "quantity": 0}],
    }
    assert (
        client.post(
            "/model-demo/orders", headers=auth["owner"], json=request
        ).status_code
        == 422
    )
    _, _, rec = ready(client, auth["owner"])
    order = draft(client, auth["owner"], rec)
    url = f"/model-demo/orders/{order['id']}"
    for lines in (
        [{"product_id": rec["rows"][1]["product_id"], "quantity": 1}],
        [{"product_id": rec["rows"][0]["product_id"], "quantity": 1}] * 2,
    ):
        assert (
            client.put(
                url, headers=auth["owner"], json={"version": 0, "lines": lines}
            ).status_code
            == 422
        )
    response = client.put(
        url,
        headers=auth["owner"],
        json={
            "version": 0,
            "lines": [
                {
                    "product_id": rec["rows"][0]["product_id"],
                    "quantity": rec["rows"][0]["quantity"],
                }
            ],
        },
    )
    assert response.status_code == 200
    assert (
        response.json()["lines"][0]["scenario_results"]
        == rec["rows"][0]["scenario_results"]
    )
