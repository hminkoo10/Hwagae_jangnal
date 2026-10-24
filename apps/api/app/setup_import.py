"""Preview and atomically import a merchant's initial workbook into an empty store."""

import io
import zipfile
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from pydantic import ValidationError
from fastapi import Depends, File, Form, HTTPException, UploadFile

from app.database import encode, now_iso, store_today
from app.imports import MAX_BYTES, MAX_ROWS
from app.schemas import ProductInput, Sale, Inventory
from app.walmart_scenarios import scenario_history

PRODUCT_COLUMNS = ["id", "name", "category", "supplier", "price", "cost", "lead_days", "shelf_days", "review_days"]
SALES_COLUMNS = ["date", "product_id", "units_sold"]
STOCK_COLUMNS = ["product_id", "measured_on", "quantity", "expires_on", "arrives_on"]


def read_workbook(filename, content, cutoff):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("상품·판매·재고를 담은 XLSX 파일을 선택하세요.")
    if len(content) > MAX_BYTES:
        raise ValueError("파일은 5MB까지 지원합니다.")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
            raise ValueError("압축 해제 크기가 30MB를 넘습니다.")
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        def records(sheet, columns):
            if sheet not in book.sheetnames:
                raise ValueError(f"'{sheet}' 시트가 필요합니다.")
            rows = book[sheet].iter_rows(values_only=True)
            header = list(next(rows, ()))
            if header != columns:
                raise ValueError(f"{sheet} 시트의 열: {', '.join(columns)}")
            result = []
            for number, row in enumerate(rows, 2):
                if number > MAX_ROWS + 1:
                    raise ValueError("시트당 10,000행까지 지원합니다.")
                if all(value is None for value in row):
                    continue
                if any(isinstance(v, str) and v.startswith("=") for v in row):
                    raise ValueError(f"{sheet} {number}행: 수식 대신 값을 입력하세요.")
                values = [v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date) else v for v in row]
                result.append(dict(zip(columns, values)))
            return result

        products = [ProductInput.model_validate(row).model_dump(mode="json") for row in records("상품", PRODUCT_COLUMNS)]
        if not 1 <= len(products) <= 50 or len({p["id"] for p in products}) != len(products):
            raise ValueError("상품은 중복 코드 없이 1~50개를 입력하세요.")
        ids = {p["id"] for p in products}
        sales = [Sale.model_validate(row).model_dump(mode="json") for row in records("판매", SALES_COLUMNS)]
        if not sales or len({(s["date"], s["product_id"]) for s in sales}) != len(sales):
            raise ValueError("판매 기록이 없거나 같은 날짜·상품이 중복되었습니다.")
        if any(s["product_id"] not in ids or s["date"] > cutoff.isoformat() for s in sales):
            raise ValueError("판매 기록의 미등록 상품 또는 미래 날짜를 확인하세요.")
        stock = {pid: {"batches": [], "incoming": [], "incoming_confirmed": True} for pid in ids}
        measured = set()
        for row in records("재고", STOCK_COLUMNS):
            pid = row["product_id"]
            if pid not in ids or row["measured_on"] != cutoff.isoformat():
                raise ValueError(f"재고는 매장 현지 오늘({cutoff}) 기준으로 입력하고 상품코드를 확인하세요.")
            measured.add(pid)
            target = "incoming" if row["arrives_on"] else "batches"
            batch = {"quantity": row["quantity"], "expires_on": row["expires_on"]}
            if target == "incoming":
                batch["arrives_on"] = row["arrives_on"]
                if date.fromisoformat(row["arrives_on"]) <= cutoff:
                    raise ValueError("예정 입고는 내일부터 입력하세요.")
            stock[pid][target].append(batch)
        if measured != ids:
            raise ValueError("모든 상품의 재고를 입력하세요. 재고가 없으면 수량 0으로 입력하세요.")
        stock = {pid: Inventory.model_validate(value).model_dump(mode="json") for pid, value in stock.items()}

        # Only this published, complete history may use the saved Walmart identities.
        # No arbitrary merchant IDs are remapped to trained embeddings.
        binding = None
        if "모델연결" in book.sheetnames:
            config = records("모델연결", ["profile", "anchor_date"])
            if len(config) != 1 or config[0]["anchor_date"] != cutoff.isoformat():
                raise ValueError("모델 연결 기준일은 매장 현지 오늘과 같아야 합니다.")
            source = scenario_history()
            sample = next((s for s in source["stores"] if s["id"] == config[0]["profile"]), None)
            if sample is None or ids != {p["id"] for p in sample["products"]}:
                raise ValueError("학습 상품과 업로드 상품이 일치하지 않습니다.")
            delta = cutoff - date.fromisoformat(source["source_end"])
            expected = {(date.fromisoformat(s["date"]) + delta).isoformat() + "/" + s["product_id"]: s["units_sold"] for s in sample["sales"]}
            actual = {s["date"] + "/" + s["product_id"]: s["units_sold"] for s in sales}
            if actual != expected:
                raise ValueError("모델 연결 파일의 판매 기록이 원본과 다릅니다. 일반 자료는 모델연결 시트 없이 가져오세요.")
            binding = (sample["id"], cutoff.isoformat(), source["source_end"], encode(sample))
        return {"products": products, "sales": sales, "inventory": stock, "binding": binding}
    except ValidationError as exc:
        raise ValueError("입력값을 확인하세요: " + exc.errors(include_url=False)[0]["msg"]) from None
    finally:
        book.close()


def empty_store(db, store_id):
    if db.execute("SELECT 1 FROM products WHERE store_id=? LIMIT 1", (store_id,)).fetchone():
        raise ValueError("초기 자료는 상품이 없는 새 매장에만 가져올 수 있습니다. 기존 매장은 판매 기록 입력을 이용하세요.")


def persist(database, store, data):
    stamp = now_iso()
    with database.connect() as db:
        db.execute("BEGIN IMMEDIATE")
        empty_store(db, store["id"])
        # Date validation is repeated after parsing, immediately before commit.
        cutoff = store_today(store)
        if data["cutoff"] != cutoff.isoformat():
            raise ValueError("매장 날짜가 바뀌었습니다. 현재 날짜 자료를 다시 확인하세요.")
        db.executemany("INSERT INTO products VALUES (?,?,?,?,?,?,?,?,?,?)", [(store["id"], *[p[key] for key in PRODUCT_COLUMNS]) for p in data["products"]])
        db.executemany("INSERT INTO sales VALUES (?,?,?,?,?,?)", [(store["id"], s["product_id"], s["date"], s["units_sold"], "file_import", stamp) for s in data["sales"]])
        db.executemany("INSERT INTO inventory VALUES (?,?,?,?,?)", [(store["id"], pid, encode(value), cutoff.isoformat(), stamp) for pid, value in data["inventory"].items()])
        if data["binding"]:
            db.execute("INSERT INTO walmart_scenarios VALUES (?,?,?,?,?)", (store["id"], *data["binding"]))
            note = "업로드 자료: 공개 Walmart 판매 날짜 이동, 상품명·가격·재고는 가정값. 매장 위치는 사용자가 지정한 날씨 위치이며 원본 Walmart 좌표가 아닙니다. 기존 체크포인트에 예보를 입력한 실험으로 정확도는 별도 검증이 필요합니다."
            db.execute("UPDATE stores SET provenance=? WHERE id=?", (note, store["id"]))
        db.execute("INSERT INTO store_forecast_settings VALUES (?,?,?) ON CONFLICT(store_id) DO UPDATE SET provider=excluded.provider,updated_at=excluded.updated_at", (store["id"], "merchant_tft", stamp))
        db.execute("UPDATE stores SET revision=revision+1,sales_revision=sales_revision+1 WHERE id=?", (store["id"],))


def mount_setup_routes(application, database, access, authenticated):
    async def inspect(store_id, file, user):
        store = access(store_id, user, owner=True)
        try:
            with database.connect() as db:
                empty_store(db, store_id)
            cutoff = store_today(store)
            data = read_workbook(file.filename or "", await file.read(MAX_BYTES + 1), cutoff)
            if data["binding"] and store["weather_provider"] != "noaa_gfs":
                raise ValueError("이 모델 연결 자료는 미국 날씨 지역이 필요합니다. 매장 지역을 댈러스·뉴욕·시애틀 중에서 선택하세요.")
            data["cutoff"] = cutoff.isoformat()
            return store, data
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from None
        except Exception:
            raise HTTPException(400, "파일을 읽을 수 없습니다. XLSX 양식과 입력값을 확인하세요.") from None
        finally:
            await file.close()

    @application.post("/stores/{store_id}/initial-data/preview", tags=["imports"])
    async def preview_initial(store_id: str, file: UploadFile = File(...), user=Depends(authenticated)):
        _, data = await inspect(store_id, file, user)
        days = sorted({s["date"] for s in data["sales"]})
        return {"product_count": len(data["products"]), "sales_count": len(data["sales"]),
                "inventory_count": len(data["inventory"]), "first_date": days[0], "last_date": days[-1],
                "measured_on": data["cutoff"], "products": data["products"], "inventory": data["inventory"],
                "model": "Walmart TFT" if data["binding"] else "날씨 반영 TFT",
                "can_save": True}

    @application.post("/stores/{store_id}/initial-data", status_code=201, tags=["imports"])
    async def import_initial(store_id: str, file: UploadFile = File(...), confirmed: bool = Form(False), user=Depends(authenticated)):
        if not confirmed:
            await file.close()
            raise HTTPException(422, "재고와 입고 예정 목록을 확인한 뒤 저장하세요.")
        store, data = await inspect(store_id, file, user)
        try:
            persist(database, store, data)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from None
        return {"saved": len(data["sales"]), "product_count": len(data["products"])}
