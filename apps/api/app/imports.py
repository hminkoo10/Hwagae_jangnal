"""Bounded, preview-first CSV/XLSX parsing without executing formulas."""

import csv
import io
import re
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from app.database import today

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
HEADERS = {
    "date": "date",
    "판매일": "date",
    "날짜": "date",
    "product_id": "product_id",
    "상품코드": "product_id",
    "상품 ID": "product_id",
    "units_sold": "units_sold",
    "판매수량": "units_sold",
    "판매 개수": "units_sold",
}


def parse_file(filename: str, content: bytes):
    if len(content) > MAX_BYTES:
        raise ValueError("파일은 5MB까지 지원합니다.")
    if filename.lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp949")
        rows = list(csv.reader(io.StringIO(text)))
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("압축 해제 크기가 30MB를 넘습니다.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        try:
            rows = []
            for row in book.active.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("한 번에 10,000행까지 지원합니다.")
        finally:
            book.close()
    else:
        raise ValueError("CSV 또는 XLSX 파일을 선택하세요.")
    if not rows:
        raise ValueError("빈 파일입니다.")
    headers = [HEADERS.get(str(value).strip(), str(value).strip()) for value in rows[0]]
    if any(headers.count(key) != 1 for key in ("date", "product_id", "units_sold")):
        raise ValueError("필수 열: date, product_id, units_sold (각각 한 번씩)")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError("한 번에 10,000행까지 지원합니다.")
    valid, errors = [], []
    seen = set()
    for number, values in enumerate(rows[1:], 2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        try:
            item = dict(zip(headers, values))
            raw_day = item.get("date")
            if isinstance(raw_day, (date, datetime)):
                day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
            else:
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(raw_day)):
                    raise ValueError("날짜는 YYYY-MM-DD 형식이어야 합니다.")
                day = date.fromisoformat(str(raw_day))
            if day > today():
                raise ValueError("미래 날짜의 판매 기록은 저장할 수 없습니다.")
            product_id = str(item.get("product_id") or "").strip()
            if not product_id or len(product_id) > 80:
                raise ValueError("상품코드를 확인하세요.")
            quantity = str(
                item.get("units_sold") if item.get("units_sold") is not None else ""
            ).strip()
            if (
                not re.fullmatch(r"\d+(?:\.0+)?", quantity)
                or float(quantity) > 1_000_000
            ):
                raise ValueError("판매수량은 0~1,000,000 사이의 정수여야 합니다.")
            key = (day.isoformat(), product_id)
            if key in seen:
                raise ValueError("파일 안에 같은 날짜·상품이 중복되었습니다.")
            seen.add(key)
            valid.append(
                {
                    "date": day.isoformat(),
                    "product_id": product_id,
                    "units_sold": int(float(quantity)),
                    "row_number": number,
                }
            )
        except (ValueError, TypeError) as exc:
            errors.append({"row": number, "message": str(exc)})
    return valid, errors
