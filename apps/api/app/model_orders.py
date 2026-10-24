"""Personal historical TFT plans. Never mutate merchant or model input records."""

import csv
import io
import json
import logging
import re
import uuid
from datetime import date, timedelta
from typing import Literal

from fastapi import Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.database import encode, now_iso
from app.model_demo import DEFAULT_SAMPLE_ID, SampleId, sample_context
from app.model_order_schemas import (
    ModelOrderConfirm,
    ModelOrderUpdate,
    ModelRecommendationRequest,
    PlanningUpdate,
)
from app.ordering import recommend
from app.schemas import OrderRequest

logger = logging.getLogger(__name__)
ASSUMPTIONS = (
    "과거 기준일의 개인 발주 시연이며 실제 공급사에 전송하지 않습니다. "
    "가격·원가·공급사·입고기간·판매기간은 시연용 가정이며 재고는 사용자가 입력합니다. "
    "검증된 TFT P10·P50·P90만 사용하고 실제 판매 정답과 기준선은 계산에서 제외합니다. "
    "세 수요 경로를 같은 비중으로 비교한 근사 계획으로 품절 확률·실제 절감액이 아닙니다. "
    "먼저 만료되는 재고부터 판매하며 기간 말 미만료 재고는 원가로 이월합니다."
)


def save_forecast_run(database, user, result):
    """Persist an allowlisted snapshot after replay has validated model output."""
    payload = {
        "sample_id": result["sample_id"],
        "as_of": result["as_of"],
        "model": {
            key: result["model"][key] for key in ("name", "version", "uses_weather")
        },
        "products": [
            {
                "product_id": product["product_id"],
                "status": product["status"],
                "reason": product["reason"],
                "points": [
                    {key: point[key] for key in ("date", "p10", "p50", "p90")}
                    for point in product["points"]
                ],
            }
            for product in result["products"]
        ],
    }
    run_id = str(uuid.uuid4())
    with database.connect() as db:
        db.execute(
            "INSERT INTO model_demo_runs VALUES (?,?,?,?,?,?)",
            (
                run_id,
                user["id"],
                result["sample_id"],
                result["as_of"],
                encode(payload),
                now_iso(),
            ),
        )
    return run_id


def _context(sample_id):
    try:
        return sample_context(sample_id)
    except Exception:
        logger.exception("TFT planning sample could not be read")
        raise HTTPException(503, "모델 체험용 자료를 준비하지 못했습니다.") from None


def _owner(db, user_id):
    return (
        db.execute(
            "SELECT 1 FROM memberships WHERE user_id=? AND role='owner' LIMIT 1",
            (user_id,),
        ).fetchone()
        is not None
    )


def _owned(db, table, record_id, user_id):
    assert table in {
        "model_demo_runs",
        "model_demo_recommendations",
        "model_demo_orders",
    }
    row = db.execute(
        f"SELECT * FROM {table} WHERE id=? AND user_id=?", (record_id, user_id)
    ).fetchone()
    if row is None:
        raise HTTPException(404, "개인 시연 결과를 찾을 수 없습니다.")
    return row


def _latest(db, user_id, sample_id):
    return db.execute(
        "SELECT * FROM model_demo_runs WHERE user_id=? AND sample_id=? ORDER BY rowid DESC LIMIT 1",
        (user_id, sample_id),
    ).fetchone()


def _workspace(db, user_id, context):
    row = db.execute(
        "SELECT revision,payload FROM model_demo_workspaces WHERE user_id=? AND sample_id=?",
        (user_id, context["store_id"]),
    ).fetchone()
    settings = (
        json.loads(row["payload"])
        if row
        else [
            {
                "product_id": product["id"],
                "price": 4000,
                "cost": 2000,
                "lead_days": 1,
                "shelf_days": 3,
                "review_days": 1,
                "stock": None,
            }
            for product in context["products"]
        ]
    )
    by_id = {item["product_id"]: item for item in settings}
    if set(by_id) != {product["id"] for product in context["products"]}:
        raise HTTPException(
            409, "시연 상품 구성이 변경되었습니다. 설정을 다시 확인하세요."
        )
    items = [
        {**by_id[product["id"]], "name": product["name"], "supplier": "시연 공급사"}
        for product in context["products"]
    ]
    return (row["revision"] if row else 0), items


def _planning(db, user_id, context):
    revision, items = _workspace(db, user_id, context)
    for item in items:
        if item["stock"] is not None:
            item["stock"] = {
                key: item["stock"][key]
                for key in ("batches", "incoming", "incoming_confirmed")
            }
    latest = _latest(db, user_id, context["store_id"])
    return {
        "sample_id": context["store_id"],
        "as_of": context["as_of"],
        "revision": revision,
        "can_confirm_export": _owner(db, user_id),
        "latest_forecast_id": latest["id"] if latest else None,
        "latest_forecast_created_at": latest["created_at"] if latest else None,
        "model": json.loads(latest["payload"])["model"] if latest else None,
        "items": items,
        "assumptions": ASSUMPTIONS,
    }


def _fresh(db, user_id, sample_id, forecast_id, revision):
    latest = _latest(db, user_id, sample_id)
    current = db.execute(
        "SELECT revision FROM model_demo_workspaces WHERE user_id=? AND sample_id=?",
        (user_id, sample_id),
    ).fetchone()
    return bool(
        latest
        and latest["id"] == forecast_id
        and (current[0] if current else 0) == revision
    )


def _require_fresh(db, user_id, sample_id, forecast_id, revision):
    if not _fresh(db, user_id, sample_id, forecast_id, revision):
        raise HTTPException(
            409,
            "예측 또는 설정이 변경되었습니다. 최신 자료로 발주안을 다시 계산하세요.",
        )


def _order_view(db, record):
    payload = json.loads(record["payload"])
    return {
        **payload,
        "order_date": payload.get("order_date", payload["as_of"]),
        "id": record["id"],
        "sample_id": record["sample_id"],
        "status": record["status"],
        "version": record["version"],
        "created_at": record["created_at"],
        "confirmed_at": record["confirmed_at"],
        "stale": not _fresh(
            db,
            record["user_id"],
            record["sample_id"],
            record["forecast_id"],
            record["revision"],
        ),
    }


def _lines(requested, recommendation, expected=None):
    ids = [line.product_id for line in requested]
    available = {
        row["product_id"]: row
        for row in recommendation["rows"]
        if row["status"] == "ready"
    }
    if len(ids) != len(set(ids)) or not set(ids).issubset(available):
        raise HTTPException(422, "계산 가능한 상품을 중복 없이 선택하세요.")
    if expected is not None and set(ids) != expected:
        raise HTTPException(422, "저장된 발주안의 모든 상품을 정확히 입력하세요.")
    lines = []
    for item in requested:
        source = available[item.product_id]
        line = {
            **source,
            "recommended_quantity": source["quantity"],
            "quantity": item.quantity,
            "amount": item.quantity * source["cost"],
        }
        if item.quantity != source["quantity"]:
            # Scenario metrics describe the recommended quantity, not an edited order.
            line.pop("scenario_loss", None)
            line.pop("scenario_results", None)
        lines.append(line)
    return lines


def _safe(value):
    if not isinstance(value, str):
        return value
    value = re.sub(
        r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]", "", value
    )
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def _export(order, format):
    rows = [
        ["개인 시연용 발주서 · 외부 발주 전송 없음"],
        [
            "과거 기준일",
            order["as_of"],
            "시연 매장",
            order["sample_id"],
            "발주 예정일",
            order.get("order_date", order["as_of"]),
        ],
        ["모델", order["model"]["name"], "버전", order["model"]["version"]],
        ["확정 시각", order["confirmed_at"], "발주안", order["id"]],
        ["가정", order["assumptions"]],
        [],
        [
            "공급사",
            "상품코드",
            "상품명",
            "입고예정일",
            "추천수량(개)",
            "확정수량(개)",
            "수정량(개)",
            "매입원가(원)",
            "합계(원)",
        ],
    ]
    rows += [
        [
            line["supplier"],
            line["product_id"],
            line["name"],
            line["arrival_date"],
            line["recommended_quantity"],
            line["quantity"],
            line["quantity"] - line["recommended_quantity"],
            line["cost"],
            line["amount"],
        ]
        for line in order["lines"]
    ]
    rows.append(["총 매입금액", order["total"]])
    rows = [[_safe(value) for value in row] for row in rows]
    if format == "csv":
        stream = io.StringIO(newline="")
        csv.writer(stream).writerows(rows)
        content, media = (
            stream.getvalue().encode("utf-8-sig"),
            "text/csv; charset=utf-8",
        )
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "시연 발주서"
        for row in rows:
            sheet.append(row)
        for cell in sheet[7]:
            cell.fill = PatternFill("solid", fgColor="167D68")
            cell.font = Font(color="FFFFFF", bold=True)
        sheet.freeze_panes = "A8"
        for column in "ABCDEFGHI":
            sheet.column_dimensions[column].width = 22
        stream = io.BytesIO()
        book.save(stream)
        content, media = (
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    return Response(
        content,
        media_type=media,
        headers={
            "Content-Disposition": f'attachment; filename="model-demo-order-{order["id"]}.{format}"'
        },
    )


def mount_model_order_routes(application, database, authenticated):
    @application.get("/model-demo/planning", tags=["model demo"])
    def planning(sample_id: SampleId = DEFAULT_SAMPLE_ID, user=Depends(authenticated)):
        context = _context(sample_id)
        with database.connect() as db:
            db.execute("BEGIN")
            return _planning(db, user["id"], context)

    @application.put("/model-demo/planning", tags=["model demo"])
    def update_planning(
        payload: PlanningUpdate,
        sample_id: SampleId = DEFAULT_SAMPLE_ID,
        user=Depends(authenticated),
    ):
        context = _context(sample_id)
        ids = [item.product_id for item in payload.items]
        if len(ids) != len(set(ids)) or set(ids) != {
            product["id"] for product in context["products"]
        }:
            raise HTTPException(
                422, "시연 매장의 모든 상품을 중복 없이 정확히 입력하세요."
            )
        items = [item.model_dump(mode="json") for item in payload.items]
        for item in items:
            stock = item["stock"]
            if stock is not None:
                if any(
                    batch["expires_on"] <= context["as_of"]
                    for batch in stock["batches"]
                ):
                    raise HTTPException(
                        422, "현재 재고의 만료일은 과거 기준일 다음 날부터 입력하세요."
                    )
                if any(
                    batch["arrives_on"] <= context["as_of"]
                    for batch in stock["incoming"]
                ):
                    raise HTTPException(
                        422, "입고 예정일은 과거 기준일 다음 날부터 입력하세요."
                    )
                stock["measured_on"] = context["as_of"]
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            current = db.execute(
                "SELECT revision FROM model_demo_workspaces WHERE user_id=? AND sample_id=?",
                (user["id"], sample_id),
            ).fetchone()
            if (current[0] if current else 0) != payload.revision:
                raise HTTPException(
                    409, "다른 화면에서 설정을 수정했습니다. 새로 불러오세요."
                )
            db.execute(
                "INSERT INTO model_demo_workspaces VALUES (?,?,?,?,?) ON CONFLICT(user_id,sample_id) DO UPDATE SET revision=excluded.revision,payload=excluded.payload,updated_at=excluded.updated_at",
                (user["id"], sample_id, payload.revision + 1, encode(items), now_iso()),
            )
            return _planning(db, user["id"], context)

    @application.post("/model-demo/recommendations", tags=["model demo"])
    def recommendations(
        payload: ModelRecommendationRequest, user=Depends(authenticated)
    ):
        with database.connect() as db:
            db.execute("BEGIN")
            run = _owned(db, "model_demo_runs", payload.forecast_id, user["id"])
            sample_id, as_of = run["sample_id"], run["as_of"]
            _require_fresh(db, user["id"], sample_id, run["id"], payload.revision)
            context = _context(sample_id)
            if as_of != context["as_of"]:
                raise HTTPException(
                    409, "시연 기준일이 변경되었습니다. 예측을 다시 실행하세요."
                )
            cutoff = date.fromisoformat(as_of)
            order_date = payload.order_date or cutoff
            if not cutoff <= order_date <= cutoff + timedelta(days=6):
                raise HTTPException(
                    422,
                    "발주 예정일은 판매 마감 기준일부터 6일 뒤까지 선택하세요. 입고·판매기간이 7일 예측 안에 들어야 합니다.",
                )
            _, items = _workspace(db, user["id"], context)
            snapshot = json.loads(run["payload"])
        forecasts = {item["product_id"]: item for item in snapshot["products"]}
        if set(forecasts) != {item["product_id"] for item in items}:
            raise HTTPException(
                409, "예측 상품 구성이 변경되었습니다. 예측을 다시 실행하세요."
            )
        rows = [
            recommend(
                {**item, "id": item["product_id"]},
                forecasts[item["product_id"]],
                item["stock"],
                cutoff,
                order_date,
            )
            for item in items
        ]
        result = {
            "id": str(uuid.uuid4()),
            "sample_id": sample_id,
            "forecast_id": run["id"],
            "as_of": as_of,
            "order_date": order_date.isoformat(),
            "revision": payload.revision,
            "model": snapshot["model"],
            "rows": rows,
            "calculated_count": sum(row["status"] == "ready" for row in rows),
            "total": sum(row["amount"] or 0 for row in rows),
            "assumptions": ASSUMPTIONS,
        }
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            _require_fresh(db, user["id"], sample_id, run["id"], payload.revision)
            db.execute(
                "INSERT INTO model_demo_recommendations VALUES (?,?,?,?,?,?,?,?)",
                (
                    result["id"],
                    user["id"],
                    sample_id,
                    run["id"],
                    payload.revision,
                    as_of,
                    encode(result),
                    now_iso(),
                ),
            )
        return result

    @application.get("/model-demo/orders", tags=["model demo"])
    def orders(sample_id: SampleId = DEFAULT_SAMPLE_ID, user=Depends(authenticated)):
        with database.connect() as db:
            db.execute("BEGIN")
            records = db.execute(
                "SELECT * FROM model_demo_orders WHERE user_id=? AND sample_id=? ORDER BY rowid DESC",
                (user["id"], sample_id),
            ).fetchall()
            return {"orders": [_order_view(db, row) for row in records]}

    @application.post("/model-demo/orders", tags=["model demo"])
    def create_order(payload: OrderRequest, user=Depends(authenticated)):
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            record = _owned(
                db, "model_demo_recommendations", payload.recommendation_id, user["id"]
            )
            _require_fresh(
                db,
                user["id"],
                record["sample_id"],
                record["forecast_id"],
                record["revision"],
            )
            recommendation = json.loads(record["payload"])
            lines = _lines(payload.lines, recommendation)
            snapshot = {
                key: recommendation[key]
                for key in ("as_of", "forecast_id", "model", "assumptions")
            }
            snapshot["order_date"] = recommendation.get(
                "order_date", recommendation["as_of"]
            )
            snapshot.update(lines=lines, total=sum(line["amount"] for line in lines))
            order_id = str(uuid.uuid4())
            db.execute(
                "INSERT INTO model_demo_orders VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (
                    order_id,
                    user["id"],
                    record["sample_id"],
                    record["id"],
                    record["forecast_id"],
                    record["revision"],
                    0,
                    "draft",
                    encode(snapshot),
                    now_iso(),
                    None,
                ),
            )
            return _order_view(
                db, _owned(db, "model_demo_orders", order_id, user["id"])
            )

    @application.put("/model-demo/orders/{order_id}", tags=["model demo"])
    def update_order(
        order_id: str, payload: ModelOrderUpdate, user=Depends(authenticated)
    ):
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            record = _owned(db, "model_demo_orders", order_id, user["id"])
            if record["status"] != "draft" or record["version"] != payload.version:
                raise HTTPException(409, "수정 가능한 최신 초안을 다시 불러오세요.")
            _require_fresh(
                db,
                user["id"],
                record["sample_id"],
                record["forecast_id"],
                record["revision"],
            )
            snapshot = json.loads(record["payload"])
            recommendation = json.loads(
                _owned(
                    db,
                    "model_demo_recommendations",
                    record["recommendation_id"],
                    user["id"],
                )["payload"]
            )
            snapshot["lines"] = _lines(
                payload.lines,
                recommendation,
                {line["product_id"] for line in snapshot["lines"]},
            )
            snapshot["total"] = sum(line["amount"] for line in snapshot["lines"])
            db.execute(
                "UPDATE model_demo_orders SET payload=?,version=version+1 WHERE id=?",
                (encode(snapshot), order_id),
            )
            return _order_view(
                db, _owned(db, "model_demo_orders", order_id, user["id"])
            )

    @application.post("/model-demo/orders/{order_id}/confirm", tags=["model demo"])
    def confirm_order(
        order_id: str, payload: ModelOrderConfirm, user=Depends(authenticated)
    ):
        with database.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            record = _owned(db, "model_demo_orders", order_id, user["id"])
            if not _owner(db, user["id"]):
                raise HTTPException(403, "점주만 시연 발주안을 확정할 수 있습니다.")
            if record["status"] != "draft" or record["version"] != payload.version:
                raise HTTPException(409, "확정 가능한 최신 초안을 다시 불러오세요.")
            _require_fresh(
                db,
                user["id"],
                record["sample_id"],
                record["forecast_id"],
                record["revision"],
            )
            db.execute(
                "UPDATE model_demo_orders SET status='confirmed',confirmed_at=?,version=version+1 WHERE id=?",
                (now_iso(), order_id),
            )
            return _order_view(
                db, _owned(db, "model_demo_orders", order_id, user["id"])
            )

    @application.get("/model-demo/orders/{order_id}/export", tags=["model demo"])
    def export_order(
        order_id: str,
        format: Literal["csv", "xlsx"] = "csv",
        user=Depends(authenticated),
    ):
        with database.connect() as db:
            db.execute("BEGIN")
            record = _owned(db, "model_demo_orders", order_id, user["id"])
            if not _owner(db, user["id"]):
                raise HTTPException(403, "점주만 시연 발주안을 내보낼 수 있습니다.")
            if record["status"] != "confirmed":
                raise HTTPException(409, "확정한 시연 발주안만 내보낼 수 있습니다.")
            order = _order_view(db, record)
        return _export(order, format)
